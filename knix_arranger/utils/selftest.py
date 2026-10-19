"""
Selbsttest des ausgelieferten Bundles (KNiX_Arranger.exe --selftest)

Prüft ohne Lizenz, ob alles, was erst bei Bedarf geladen wird, im
kompilierten Bundle vorhanden ist: Datendateien, Firmenschrift,
Qt-Übersetzung, PDF-Erzeugung (PyMuPDF), Kryptografie und Excel-Import --
und ob sich das Hauptfenster aufbauen lässt (ohne es anzuzeigen).
Der Release-Workflow ruft ihn nach jedem Build auf -- ein fehlendes Modul
fällt so vor der Auslieferung auf und nicht erst beim Betatester.

Rückgabe: 0 = alles in Ordnung, 1 = mindestens ein Fehler (Details im Log).
"""
from __future__ import annotations
import json
import logging
import os
import tempfile

logger = logging.getLogger("knix_arranger.selftest")

_PACKAGE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def _check_data_files() -> None:
    for sub in ("config", "i18n", "data"):
        folder = os.path.join(_PACKAGE_DIR, sub)
        names = [n for n in os.listdir(folder) if n.endswith(".json")]
        if not names:
            raise RuntimeError(f"keine JSON-Dateien in {sub}/")
        for name in names:
            with open(os.path.join(folder, name), encoding="utf-8") as f:
                json.load(f)
    handbook = os.path.join(_PACKAGE_DIR, "data", "KNiX_Arranger_Handbuch.pdf")
    if not os.path.exists(handbook):
        raise RuntimeError("Bedienungsanleitung fehlt")


def _check_fonts() -> None:
    from .fonts import _BUNDLED_DIR
    for name in ("Inter-Regular.ttf", "Inter-Bold.ttf"):
        if not os.path.exists(os.path.join(_BUNDLED_DIR, name)):
            raise RuntimeError(f"{name} fehlt")


def _check_qt_translation() -> None:
    from PySide6.QtCore import QTranslator, QLibraryInfo
    path = QLibraryInfo.path(QLibraryInfo.LibraryPath.TranslationsPath)
    if not QTranslator().load("qtbase_de", path):
        raise RuntimeError(f"qtbase_de nicht gefunden in {path}")


def _check_pdf() -> None:
    import fitz
    from .pdf_generator import PdfGenerator
    fd, path = tempfile.mkstemp(suffix=".pdf")
    os.close(fd)
    try:
        pdf = PdfGenerator(title="Selbsttest")
        pdf.add_heading("Selbsttest", level=2)
        pdf.add_paragraph("Umlaute: äöü ÄÖÜ – Schweizer ss")
        pdf.save(path)
        with fitz.open(path) as doc:
            if "Selbsttest" not in doc[0].get_text():
                raise RuntimeError("Text im PDF nicht lesbar")
    finally:
        os.remove(path)


def _check_crypto() -> None:
    from cryptography.hazmat.primitives import serialization
    from ..services.license_service import _PUBLIC_KEY_PEM
    serialization.load_pem_public_key(_PUBLIC_KEY_PEM)
    import pyzipper  # noqa: F401  (ETS-Projekte mit Passwort)


def _check_excel() -> None:
    import openpyxl
    openpyxl.Workbook()


def _check_qr() -> None:
    """QR-Erkennung der Secure-Zertifikate (zxing-cpp) und Kamera-Modul:
    beides wird erst im Dialog «Zertifikate einlesen» geladen."""
    from PySide6.QtGui import QImage
    from PySide6.QtMultimedia import QMediaDevices  # noqa: F401
    from PySide6.QtMultimediaWidgets import QVideoWidget  # noqa: F401
    from ..services.knx_secure_certificate import read_qr_texts
    image = QImage(64, 64, QImage.Format_Grayscale8)
    image.fill(255)
    if read_qr_texts(image):
        raise RuntimeError("QR-Code in leerem Bild erkannt")


def _check_main_window() -> None:
    """Hauptfenster aufbauen (ohne Anzeigen und ohne Startdialoge): 1.2.5
    stürzte mit Nuitka 4.3 schon dabei ab (Segmentation Fault), der übrige
    Selbsttest lief durch. Ein Absturz beendet den Prozess mit Fehlercode."""
    from PySide6.QtWidgets import QApplication
    from ..ui.main_window import MainWindow
    app = QApplication.instance() or QApplication([])
    window = MainWindow()
    window.deleteLater()
    del app


CHECKS = [
    ("Datendateien", _check_data_files),
    ("Firmenschrift", _check_fonts),
    ("Qt-Übersetzung", _check_qt_translation),
    ("PDF-Erzeugung", _check_pdf),
    ("Kryptografie", _check_crypto),
    ("Excel", _check_excel),
    ("QR-Code und Kamera", _check_qr),
    ("Hauptfenster", _check_main_window),
]


def run_selftest() -> int:
    failed = 0
    for name, check in CHECKS:
        try:
            check()
            logger.info(f"Selbsttest {name}: OK")
        except Exception as e:
            failed += 1
            logger.error(f"Selbsttest {name}: FEHLER – {e}", exc_info=True)
    logger.info("Selbsttest bestanden." if not failed
                else f"Selbsttest: {failed} Fehler.")
    return 1 if failed else 0
