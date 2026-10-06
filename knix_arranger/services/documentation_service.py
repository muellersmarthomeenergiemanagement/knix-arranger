"""
Dokumentations-Service (FA-1900, FA-2000, FA-2100)
Erzeugt Inbetriebnahme-Checklisten, Abnahmeprotokolle,
Bedienungsanleitungen und Revisionspakete.
"""
from __future__ import annotations
import logging
import os
from datetime import datetime

from .building_service import BuildingService
from ..models.project import KnxProject
from ..models.building import Room
from ..models.documentation import (
    CommissioningChecklist, ChecklistItem, AcceptanceProtocol, Defect,
    ITEM_KIND_DEVICE, ITEM_KIND_FUNCTION, RESULT_OPEN,
)
from ..utils.pdf_generator import PdfGenerator
from .report_sorting import physical_address_key, sorted_rooms
from ..utils.excel_generator import ExcelGenerator, HAS_OPENPYXL

logger = logging.getLogger("knix_arranger.documentation")


# Standard-Prüfpunkte für Inbetriebnahme (FA-1902)
STANDARD_CHECK_ITEMS = [
    ("Montage", "Gerät korrekt montiert und beschriftet"),
    ("Busverbindung", "KNX-Busverbindung hergestellt und geprueft"),
    ("Physikalische Adresse", "Physikalische Adresse programmiert"),
    ("Applikation", "Applikationsprogramm geladen"),
    ("Funktion Schalten", "Schaltfunktion E/A getestet"),
    ("Funktion Dimmen", "Dimmfunktion getestet (falls vorhanden)"),
    ("Funktion Jalousie", "Jalousie AUF/AB getestet (falls vorhanden)"),
    ("Funktion Heizung", "Heizungsregelung getestet (falls vorhanden)"),
    ("Rueckmeldung", "Rückmeldungen (LED/Status) prüfen"),
    ("Szenen", "Szenen-Aufrufe getestet (falls vorhanden)"),
]



def _period_text(sp) -> str:
    """Datumsbereich eines Schaltzeitpunkts als TT.MM.JJJJ – TT.MM.JJJJ."""
    def fmt(value: str) -> str:
        try:
            return datetime.strptime(value, "%Y-%m-%d").strftime("%d.%m.%Y")
        except ValueError:
            return value or "…"
    if not (sp.date_range_start or sp.date_range_end):
        return "ganzjährig"
    return f"{fmt(sp.date_range_start)} – {fmt(sp.date_range_end)}"

class DocumentationService:
    """Erzeugt Dokumentation für KNX-Projekte."""

    def __init__(self, project: KnxProject, company_profile=None):
        self.project = project
        self._company_profile = company_profile

    def _make_pdf(self, title: str) -> PdfGenerator:
        """Erstellt PdfGenerator mit Firmenprofil und Projektdaten."""
        return PdfGenerator(
            title=title,
            company_profile=self._company_profile,
            project_info=self.project.project_info,
        )

    # -- Inbetriebnahme-Checkliste (FA-1901) --

    def create_checklists(self) -> list[CommissioningChecklist]:
        """Erzeugt Inbetriebnahme-Checklisten für alle Räume (FA-1901) +
        DALI-Notlicht-Prüfpunkte (FA-2804).

        Liefert strukturierte Items (ITEM_KIND_DEVICE + ITEM_KIND_FUNCTION)
        mit be_id-Verknüpfung für die CommissioningView.
        """
        all_rooms = self._checklist_rooms()
        checklists = []
        for room in all_rooms:
            checklist = self._create_room_checklist(room)
            checklists.append(checklist)

        # DALI-Notlicht-Prüfpunkte anhängen (FA-2804)
        if self.project.dali_configs:
            from .dali_service import DaliService
            svc = DaliService()
            dali_checklist = CommissioningChecklist(
                room_id="__dali__",
                room_name="DALI Notbeleuchtung",
                date=datetime.now().strftime("%Y-%m-%d"),
            )
            for gw in self.project.dali_configs.values():
                for item_dict in svc.generate_emergency_checklist_items(gw):
                    dali_checklist.items.append(ChecklistItem(
                        item_kind=ITEM_KIND_DEVICE,
                        check_type=item_dict["category"],
                        description=item_dict["text"],
                        room_id="__dali__",
                        gewerk_code="LDA",
                    ))
            if dali_checklist.items:
                checklists.append(dali_checklist)

        # Verteiler-Geräte (Aktoren, Gateways, Koppler, Netzteile) anhängen
        checklists.extend(self._create_verteiler_checklists())

        return checklists

    # Gerätetypen mit physikalischer Adresse, die im Verteiler sitzen und
    # nicht über room.bedienelemente erfasst werden (Aktoren sind keine BEs).
    _VERTEILER_DEVICE_TYPES = ("actor", "gateway", "coupler", "power_supply")

    def _verteiler_devices_by_location(self) -> dict[str, list]:
        """Gruppiert alle Verteiler-Geräte (Aktoren/Gateways/Koppler/Netzteile)
        nach Einbauort (UV/HV), sortiert nach Ortsname."""
        by_location: dict[str, list] = {}
        for area in self.project.topology.areas:
            for line in area.lines:
                for device in line.devices:
                    if device.device_type not in self._VERTEILER_DEVICE_TYPES:
                        continue
                    location = device.installation_location or "Nicht zugeordnet"
                    by_location.setdefault(location, []).append(device)
        return dict(sorted(by_location.items()))

    def _create_verteiler_checklists(self) -> list[CommissioningChecklist]:
        """Geräte-Grundprüfungen für Verteiler-Geräte, gruppiert nach UV/HV.

        Aktoren, Gateways, Koppler und Netzteile sitzen im Verteiler statt im
        Raum und tauchen daher nicht in room.bedienelemente auf – sie haben
        aber ebenfalls eine physikalische Adresse und ein Applikationsprogramm,
        die bei der Inbetriebnahme geprüft werden müssen.
        """
        checklists = []
        for location, devices in self._verteiler_devices_by_location().items():
            room_id = f"__verteiler__{location}"
            checklist = CommissioningChecklist(
                room_id=room_id,
                room_name=f"Verteiler {location}",
                date=datetime.now().strftime("%Y-%m-%d"),
            )
            for idx, device in enumerate(devices):
                dev_number = device.physical_address or f"#{idx + 1}"
                for check_type, description in self._DEVICE_CHECKS:
                    checklist.items.append(ChecklistItem(
                        item_kind=ITEM_KIND_DEVICE,
                        check_type=check_type,
                        description=description,
                        room_id=room_id,
                        be_type=device.product or device.device_type,
                        be_number=dev_number,
                        gewerk_code=device.device_type,
                    ))
            if checklist.items:
                checklists.append(checklist)
        return checklists

    def _create_room_checklist(self, room: Room) -> CommissioningChecklist:
        """Erzeugt eine vollständige Checkliste für einen Raum.

        Pro aktivem Bedienelement:
        - 4 Geräte-Grundprüfungen (item_kind=device)
        - Je eine Funktionsprüfung pro GA-Zuordnung (item_kind=function)
        """
        checklist = CommissioningChecklist(
            room_id=room.id,
            room_name=f"{room.number} {room.name}",
            date=datetime.now().strftime("%Y-%m-%d"),
        )
        active_bes = [be for be in room.bedienelemente if not be.suppressed]
        for be in active_bes:
            be_num = be.participant_number or ""
            # Geräte-Grundprüfungen
            for check_type, description in self._DEVICE_CHECKS:
                checklist.items.append(ChecklistItem(
                    item_kind=ITEM_KIND_DEVICE,
                    check_type=check_type,
                    description=description,
                    room_id=room.id,
                    be_id=be.id,
                    be_type=be.element_type,
                    be_number=be_num,
                ))
            # GA-Funktionsprüfungen
            for fa in be.function_assignments:
                checklist.items.append(ChecklistItem(
                    item_kind=ITEM_KIND_FUNCTION,
                    check_type=fa.button_channel,
                    description=fa.description,
                    function_ga=fa.function_ga,
                    room_id=room.id,
                    be_id=be.id,
                    be_type=be.element_type,
                    be_number=be_num,
                ))
        return checklist

    def init_project_checklists(self) -> int:
        """Initialisiert project.checklists einmalig aus aktuellem Projektstand.

        Überschreibt nichts, falls bereits Checklisten vorhanden sind.
        Gibt Anzahl generierter Items zurück.
        """
        if self.project.checklists:
            return 0
        self.project.checklists = self.create_checklists()
        return sum(len(cl.items) for cl in self.project.checklists)

    def sync_project_checklists(self) -> tuple[int, int]:
        """Gleicht project.checklists mit aktuellem Projektstand ab.

        Neue Räume/BEs/GAs werden hinzugefügt.
        Bestehende Items mit Ergebnissen bleiben unverändert.
        Gibt (hinzugefügt, gesamt) zurück.
        """
        fresh = self.create_checklists()
        added = 0

        for fresh_cl in fresh:
            existing_cl = next(
                (c for c in self.project.checklists if c.room_id == fresh_cl.room_id),
                None,
            )
            if existing_cl is None:
                self.project.checklists.append(fresh_cl)
                added += len(fresh_cl.items)
                continue

            existing_cl.room_name = fresh_cl.room_name
            existing_keys = {
                (i.be_type, i.be_number, i.check_type, i.function_ga)
                for i in existing_cl.items
            }
            # be_id-Lookup: vorhandene UUID für (be_type, be_number) übernehmen,
            # da auto_assign_functions bei jedem Aufruf neue UUIDs vergibt.
            known_be_ids: dict[tuple, str] = {
                (i.be_type, i.be_number): i.be_id
                for i in existing_cl.items
                if i.be_id
            }
            for item in fresh_cl.items:
                key = (item.be_type, item.be_number, item.check_type, item.function_ga)
                if key not in existing_keys:
                    be_key = (item.be_type, item.be_number)
                    if be_key in known_be_ids:
                        item.be_id = known_be_ids[be_key]
                    existing_cl.items.append(item)
                    added += 1

        total = sum(len(cl.items) for cl in self.project.checklists)
        return added, total

    # Geräte-Grundprüfungen (Erfassung in der Inbetriebnahme-Ansicht je Punkt)
    _DEVICE_CHECKS = [
        ("Montage",           "Gerät montiert und beschriftet"),
        ("Physik. Adresse",   "Physikalische Adresse programmiert"),
        ("Applikation",       "Applikationsprogramm geladen"),
        ("Kommunikation",     "Bustelegramme empfangen und gesendet"),
    ]
    _DEVICE_CHECK_TEXT = ("montiert und beschriftet · physikalische Adresse · "
                          "Applikation geladen · Kommunikation")

    # Ausdruck (PDF/Excel): eine Zeile je Taste, breite Bemerkungsspalte
    _CL_HEADERS = ["Taste", "Funktion", "Gruppenadressen", "OK", "Bemerkung"]
    _CL_PDF_WIDTHS = [0.09, 0.24, 0.31, 0.06, 0.30]
    _CL_XL_WIDTHS = [13, 30, 42, 11, 44]  # A–E in Zeichen
    _XL_CENTER_COLS = [4]                 # OK-Spalte zentrieren
    _XL_WRAP_COLS = [2, 3, 5]
    _XL_CHECKBOX = "☐"

    def _checklist_rooms(self):
        """Alle Räume mit aktuellen Funktionszuordnungen – aus einer Kopie,
        damit das Erzeugen von Checklisten das Projekt nicht verändert."""
        from .sensor_service import project_for_export
        export_project = project_for_export(self.project)
        return sorted_rooms(export_project.areal)

    def _checklist_sections(self):
        """Inhalt des Ausdrucks: [(Stockwerk, Raum, [(Titel, Zeilen)])].

        Zeile = (Taste, Funktion, Gruppenadressen, Ergebnis-Schlüssel). Je
        Taste eine Zeile mit allen GAs (Befehl, LED); die Funktion wie in der
        Bedienungsanleitung (gesteuertes Objekt, eigene Bezeichnung). Die
        Schlüssel verweisen auf die je GA erfassten Ergebnisse der
        Inbetriebnahme-Ansicht (room_id, be_type, be_number, Kanal, GA)."""
        from .sensor_service import project_for_export
        from .user_manual import UserManualBuilder
        from .bedienelement_layout import group_assignments, parse_button
        from .report_service import _ga_line

        project = project_for_export(self.project)
        builder = UserManualBuilder(project, snapshot=False)
        imported = project.topology.is_imported
        floor_by_room = {}
        for building in project.areal.buildings:
            for wing in building.wings:
                for floor in wing.floors:
                    for apt in floor.apartments:
                        for room in apt.rooms:
                            floor_by_room[room.id] = BuildingService.floor_label(
                                project.areal, floor)

        sections = []
        for room in sorted_rooms(project.areal):
            cards = []
            for be in room.bedienelemente:
                if not be.is_shown(imported):
                    continue
                be_num = be.participant_number or ""
                rows = []
                dev_keys = [(room.id, be.element_type, be_num, pt, "")
                            for pt, _ in self._DEVICE_CHECKS]
                rows.append(("Gerät", "Grundprüfung", self._DEVICE_CHECK_TEXT, dev_keys))
                labels = {}
                if be.function_assignments:
                    try:
                        labels = {(kl.key.number, kl.key.side): kl.label
                                  for kl in builder.key_lines(be, room.name)}
                    except Exception:
                        labels = {}
                fa_keys: dict[tuple, list] = {}
                for fa in be.function_assignments:
                    parsed = parse_button(fa.button_channel)
                    k = ((parsed[0].number, parsed[0].side, parsed[0].variant)
                         if parsed else ("~", fa.button_channel or fa.description))
                    fa_keys.setdefault(k, []).append(
                        (room.id, be.element_type, be_num, fa.button_channel, fa.function_ga))
                for row in group_assignments(be.function_assignments, builder._resolve):
                    if row.key is None:
                        k = ("~", row.name)
                        taste, function = "–", row.name
                    else:
                        k = (row.key.number, row.key.side, row.key.variant)
                        taste = row.key.label()
                        function = labels.get((row.key.number, row.key.side), "") \
                            if not row.key.variant else ""
                        function = function or (
                            "langer Tastendruck" if row.key.variant == "lang"
                            else row.key.variant or "–")
                    gas = "\n".join(_ga_line(g) for g in row.gas)
                    if row.led_gas:
                        gas += "\n" + "\n".join("LED: " + _ga_line(g) for g in row.led_gas)
                    rows.append((taste, function, gas.strip() or "–", fa_keys.get(k, [])))
                device = builder.device_by_addr.get(be_num)
                product = (device.product_name or device.product) if device else be.product_name
                title = " · ".join(p for p in (be_num or "ohne Adresse",
                                               be.element_type or "Bedienelement", product) if p)
                cards.append((title, rows))
            if cards:
                floor = floor_by_room.get(room.id, "") or "Ohne Stockwerk"
                sections.append((floor, f"{room.number} {room.name}".strip(), cards))
        return sections

    def _saved_results(self) -> dict[tuple, ChecklistItem]:
        saved: dict[tuple, ChecklistItem] = {}
        for cl in self.project.checklists:
            for item in cl.items:
                saved[(item.room_id, item.be_type, item.be_number,
                       item.check_type, item.function_ga)] = item
        return saved

    @staticmethod
    def _combined_result(keys: list, saved: dict) -> tuple[str, str]:
        """Ergebnis einer Ausdruckzeile aus den je GA erfassten Punkten:
        ein Mangel -> Mangel, alle OK/n.a. -> OK, sonst offen."""
        items = [saved[k] for k in keys if k in saved]
        notes = "; ".join(dict.fromkeys(i.notes for i in items if i.notes))
        results = [i.result for i in items if i.result]
        if "Mangel" in results:
            return "Mangel", notes
        if items and len(results) == len(items) and all(r in ("OK", "n/a") for r in results):
            return "OK", notes
        return "", notes

    def export_checklists_pdf(self, filepath: str,
                              checklists: list[CommissioningChecklist] | None = None):
        """Inbetriebnahme-Checkliste als PDF (FA-1904): Stockwerk → Raum →
        Gerät, je Taste eine Zeile, breite Bemerkungsspalte; danach die
        Verteiler mit einer Zeile je Gerät."""
        saved = self._saved_results()
        pdf = self._make_pdf("Inbetriebnahme-Checkliste")
        pdf.add_heading("Inbetriebnahme-Checkliste", level=1)
        pdf.add_paragraph(
            f"Projekt: {self.project.name} | Datum: {datetime.now().strftime('%d.%m.%Y')}")
        pdf.add_note("Hinweis:", "Je Taste eine Zeile mit allen Gruppenadressen. "
                                 "OK ankreuzen, Mängel in der Spalte Bemerkung festhalten.")

        def mark(keys):
            result, notes = self._combined_result(keys, saved)
            return {"OK": "[x]", "Mangel": "[!]"}.get(result, "[ ]"), notes

        current_floor = None
        for floor, room_label, cards in self._checklist_sections():
            if floor != current_floor:
                pdf.add_page_break()
                pdf.add_heading(floor, level=2)
                current_floor = floor
            else:
                pdf.add_conditional_break(min_height=140)
            pdf.add_heading(room_label, level=3)
            for title, rows in cards:
                pdf.add_conditional_break(min_height=min(60 + 26 * len(rows), 400))
                pdf.add_card_header(title, "", bookmark=f"{room_label}: {title}")
                table = []
                for taste, function, gas, keys in rows:
                    ok, notes = mark(keys)
                    table.append([taste, function, gas, ok, notes])
                pdf.add_table(self._CL_HEADERS, table, col_widths=self._CL_PDF_WIDTHS)

        # DALI-Notlicht-Prüfpunkte (FA-2804)
        if self.project.dali_configs:
            from .dali_service import DaliService
            svc = DaliService()
            pdf.add_page_break()
            pdf.add_heading("DALI Notbeleuchtung", level=2)
            for gw in self.project.dali_configs.values():
                rows = [[d["category"], d["text"], "", "[ ]", ""]
                        for d in svc.generate_emergency_checklist_items(gw)]
                if rows:
                    pdf.add_table(self._CL_HEADERS, rows, col_widths=self._CL_PDF_WIDTHS)

        # Verteiler: eine Zeile je Gerät
        verteiler = self._verteiler_devices_by_location()
        if verteiler:
            pdf.add_page_break()
            pdf.add_heading("Verteiler", level=2)
            pdf.add_note("Je Gerät:", self._DEVICE_CHECK_TEXT)
            from .report_service import _clean_location
            for location, devices in verteiler.items():
                pdf.add_heading(_clean_location(location), level=3)
                table = []
                for idx, device in enumerate(devices):
                    label = device.product_name or device.product or device.device_type
                    number = device.physical_address or f"#{idx + 1}"
                    keys = [(f"__verteiler__{location}", device.product or device.device_type,
                             number, pt, "") for pt, _ in self._DEVICE_CHECKS]
                    ok, notes = mark(keys)
                    table.append([number, label, "", ok, notes])
                pdf.add_table(["Adresse", "Gerät", "", "OK", "Bemerkung"], table,
                              col_widths=[0.09, 0.45, 0.10, 0.06, 0.30])

        pdf.save(filepath)
        logger.info(f"Checklisten exportiert: {filepath}")

    # Auswahl in der Spalte OK (Excel-Auswahlliste) <-> Ergebnis in KNiX
    XL_CHOICES = {"☐ offen": "", "✓ OK": "OK", "⚠ Mangel": "Mangel", "– n/a": "n/a"}
    XL_OPEN = "☐ offen"

    @classmethod
    def _xl_choice(cls, result: str) -> str:
        return next((k for k, v in cls.XL_CHOICES.items() if v == result), cls.XL_OPEN)

    def export_checklists_excel(self, filepath: str,
                                checklists: list[CommissioningChecklist] | None = None):
        """Inbetriebnahme-Checkliste als Excel (FA-1904) zum Ausfüllen auf der
        Baustelle und Wiedereinlesen (import_checklists_excel):
        - Blatt "Übersicht" mit Fortschritt je Stockwerk (Formeln),
        - je Stockwerk ein Blatt, dazu "Verteiler" und ggf. "DALI Notlicht",
        - Spalte OK als Auswahlliste (offen / OK / Mangel / n/a) mit Farbe,
        - ausgeblendete Spalte F mit dem Bezug zu den Prüfpunkten der App."""
        if not HAS_OPENPYXL:
            raise ImportError("openpyxl wird für Excel-Export benötigt.")
        import json
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.formatting.rule import FormulaRule
        from openpyxl.styles import PatternFill, Font, Alignment
        from .report_service import _clean_location

        saved = self._saved_results()
        headers = self._CL_HEADERS + ["Bezug (nicht ändern)"]
        widths = self._CL_XL_WIDTHS + [4]

        def row_cells(keys, taste, function, gas):
            result, notes = self._combined_result(keys, saved)
            return [taste, function, gas, self._xl_choice(result), notes,
                    json.dumps([list(k) for k in keys], ensure_ascii=False)]

        excel = ExcelGenerator(title="Übersicht", project_name=self.project.name)
        sheets: list[str] = []

        def new_sheet(name: str):
            excel.add_sheet(name)
            excel.set_column_widths(widths)
            excel.set_print_options(orientation="landscape")
            sheets.append(excel._current_sheet.title)

        def table(rows):
            excel.add_table(headers, rows, center_cols=self._XL_CENTER_COLS,
                            wrap_cols=self._XL_WRAP_COLS, data_row_height=20)

        current_floor = None
        for floor, room_label, cards in self._checklist_sections():
            if floor != current_floor:
                new_sheet(floor)
                excel.add_heading(f"Inbetriebnahme-Checkliste – {floor}", level=1)
                current_floor = floor
            excel.add_heading(room_label, level=2)
            for title, rows in cards:
                excel.add_heading(title, level=3)
                table([row_cells(keys, taste, function, gas)
                       for taste, function, gas, keys in rows])
                excel.add_empty_row(height=4)

        verteiler = self._verteiler_devices_by_location()
        if verteiler:
            new_sheet("Verteiler")
            excel.add_heading("Inbetriebnahme-Checkliste – Verteiler", level=1)
            excel.add_paragraph(f"Je Gerät: {self._DEVICE_CHECK_TEXT}")
            for location, devices in verteiler.items():
                excel.add_heading(_clean_location(location), level=2)
                rows = []
                for idx, device in enumerate(devices):
                    label = device.product_name or device.product or device.device_type
                    number = device.physical_address or f"#{idx + 1}"
                    keys = [(f"__verteiler__{location}", device.product or device.device_type,
                             number, pt, "") for pt, _ in self._DEVICE_CHECKS]
                    rows.append(row_cells(keys, number, label, ""))
                table(rows)
                excel.add_empty_row(height=4)

        if self.project.dali_configs:
            from .dali_service import DaliService
            svc = DaliService()
            rows = []
            for gw in self.project.dali_configs.values():
                for d in svc.generate_emergency_checklist_items(gw):
                    keys = [("__dali__", "", "", d["category"], "")]
                    rows.append(row_cells(keys, d["category"], d["text"], ""))
            if rows:
                new_sheet("DALI Notlicht")
                excel.add_heading("DALI Notbeleuchtung", level=1)
                table(rows)

        # Auswahlliste, Farben, Bezugsspalte je Blatt
        green = PatternFill(start_color="DCEFD6", end_color="DCEFD6", fill_type="solid")
        red = PatternFill(start_color="F8D7D3", end_color="F8D7D3", fill_type="solid")
        grey = PatternFill(start_color="ECECEC", end_color="ECECEC", fill_type="solid")
        choices = ",".join(self.XL_CHOICES)
        for name in sheets:
            ws = excel.wb[name]
            last = max(ws.max_row, 2)
            dv = DataValidation(type="list", formula1=f'"{choices}"', allow_blank=True,
                                showErrorMessage=True, errorTitle="Ergebnis",
                                error="Bitte aus der Liste wählen.")
            ws.add_data_validation(dv)
            for r in range(1, last + 1):
                if ws.cell(r, 4).value in self.XL_CHOICES:
                    dv.add(ws.cell(r, 4))
            area = f"A1:E{last}"
            ws.conditional_formatting.add(area, FormulaRule(formula=['$D1="✓ OK"'], fill=green))
            ws.conditional_formatting.add(area, FormulaRule(formula=['$D1="⚠ Mangel"'], fill=red))
            ws.conditional_formatting.add(area, FormulaRule(formula=['$D1="– n/a"'], fill=grey))
            ws.column_dimensions["F"].hidden = True
            ws.print_area = area

        # Übersicht: Fortschritt je Blatt per Formel
        ws = excel.wb["Übersicht"]
        excel._current_sheet = ws
        excel._row = 1
        excel.set_column_widths([30, 12, 12, 12, 12, 12])
        excel.set_print_options(orientation="landscape")
        excel.add_header()
        excel.add_empty_row()
        excel.add_paragraph("In den Blättern je Stockwerk in der Spalte OK auswählen "
                            "(Zelle anklicken, Pfeil): offen, OK, Mangel oder n/a. "
                            "Die ausgefüllte Datei lässt sich in KNiX wieder einlesen "
                            "(Inbetriebnahme → Excel einlesen).")
        excel.add_empty_row()
        start = excel._row
        excel.add_table(["Blatt", "Zeilen", "OK", "Mangel", "n/a", "offen"],
                        [[n, "", "", "", "", ""] for n in sheets],
                        center_cols=[2, 3, 4, 5, 6])
        for i, name in enumerate(sheets):
            r = start + 1 + i
            ref = f"'{name}'!D:D"
            ws.cell(r, 3).value = f'=COUNTIF({ref},"✓ OK")'
            ws.cell(r, 4).value = f'=COUNTIF({ref},"⚠ Mangel")'
            ws.cell(r, 5).value = f'=COUNTIF({ref},"– n/a")'
            ws.cell(r, 6).value = f'=COUNTIF({ref},"☐ offen")'
            ws.cell(r, 2).value = f"=SUM(C{r}:F{r})"
        total = start + 1 + len(sheets)
        ws.cell(total, 1).value = "Total"
        ws.cell(total, 1).font = Font(name="Inter", size=9, bold=True)
        for col in range(2, 7):
            letter = "ABCDEF"[col - 1]
            c = ws.cell(total, col)
            c.value = f"=SUM({letter}{start + 1}:{letter}{total - 1})"
            c.font = Font(name="Inter", size=9, bold=True)
            c.alignment = Alignment(horizontal="center")
        excel.wb.active = 0

        excel.save(filepath)
        logger.info(f"Checklisten Excel exportiert: {filepath}")

    def import_checklists_excel(self, filepath: str) -> dict:
        """Liest eine ausgefüllte Checkliste (export_checklists_excel) ein:
        Ergebnis der Spalte OK und Bemerkung gehen auf alle Prüfpunkte der
        Zeile (Bezug in Spalte F). "offen" bzw. leer lässt vorhandene
        Ergebnisse unverändert. Fehlen die Prüfpunkte noch, werden sie aus
        dem aktuellen Projektstand angelegt.

        Gibt {"rows": gelesene Zeilen, "items": geänderte Prüfpunkte,
        "unknown": Zeilen ohne passenden Prüfpunkt} zurück."""
        import json
        from openpyxl import load_workbook

        if not self.project.checklists:
            self.init_project_checklists()
        else:
            self.sync_project_checklists()
        by_key: dict[tuple, ChecklistItem] = {}
        for cl in self.project.checklists:
            for item in cl.items:
                by_key[(item.room_id, item.be_type, item.be_number,
                        item.check_type, item.function_ga)] = item

        stats = {"rows": 0, "items": 0, "unknown": 0}
        wb = load_workbook(filepath, data_only=True)
        for ws in wb.worksheets:
            if ws.title == "Übersicht":
                continue
            for row in ws.iter_rows(min_row=1, max_col=6, values_only=True):
                ref = row[5] if len(row) > 5 else None
                if not isinstance(ref, str) or not ref.startswith("[["):
                    continue
                try:
                    keys = [tuple(k) for k in json.loads(ref)]
                except (ValueError, TypeError):
                    continue
                choice = (row[3] or "").strip() if isinstance(row[3], str) else ""
                result = self.XL_CHOICES.get(choice, "")
                notes = (row[4] or "").strip() if isinstance(row[4], str) else ""
                items = [by_key[k] for k in keys if k in by_key]
                stats["rows"] += 1
                if not items:
                    stats["unknown"] += 1
                    continue
                for item in items:
                    changed = False
                    if result and item.result != result:
                        item.result = result
                        changed = True
                    if notes and item.notes != notes:
                        item.notes = notes
                        changed = True
                    if changed:
                        stats["items"] += 1
        logger.info(f"Checklisten Excel eingelesen: {filepath} {stats}")
        return stats

    # -- Abnahmeprotokoll (FA-1911) --

    def create_acceptance_protocol(self, integrator_name: str = "",
                                   client_name: str = "",
                                   ) -> AcceptanceProtocol:
        """Erzeugt ein leeres Abnahmeprotokoll (FA-1911). Integrator und
        Bauherr ohne Angabe aus Firmen- bzw. Kundenprofil."""
        company = self._company_profile
        if not integrator_name and company:
            integrator_name = ", ".join(
                p for p in (company.user_name, company.company_name) if p)
        client = getattr(self.project, "client_profile", None)
        if not client_name and client:
            client_name = client.name
        protocol = AcceptanceProtocol(
            date=datetime.now().strftime("%d.%m.%Y"),
            integrator_name=integrator_name,
            integrator_role=company.role if company else "",
            client_name=client_name,
            checklists=self.create_checklists(),
        )
        return protocol

    def _plant_scope(self) -> list[list[str]]:
        """Umfang der Anlage für das Abnahmeprotokoll."""
        project = self.project
        imported = project.topology.is_imported
        devices = [d for a in project.topology.areas for ln in a.lines for d in ln.devices
                   if d.physical_address and not d.physical_address.endswith("-")]
        lines = sum(1 for a in project.topology.areas for ln in a.lines if ln.devices)
        bes = [be for r in project.all_rooms for be in r.bedienelemente if be.is_shown(imported)]
        operable = sum(1 for be in bes if be.is_operable)
        sensors = len(bes) - operable
        actors = sum(1 for d in devices if d.device_type in ("actor", "gateway"))
        gas = sum(1 for ga in project.group_addresses.all_addresses() if not ga.is_placeholder)
        rooms = sum(1 for r in project.all_rooms if r.bedienelemente)
        secure = project.knx_secure.enabled if project.knx_secure else False
        return [
            ["Busteilnehmer", str(len(devices)), "Linien", str(lines)],
            ["Bedienelemente", str(operable), "Sensoren", str(sensors)],
            ["Aktoren und Gateways", str(actors), "Gruppenadressen", str(gas)],
            ["Räume mit Bedienung", str(rooms), "KNX Secure", "ja" if secure else "nein"],
        ]

    def export_acceptance_protocol(self, filepath: str,
                                   protocol: AcceptanceProtocol | None = None):
        """Abnahmeprotokoll als Formular (FA-1911 bis FA-1914): Projektdaten,
        Umfang der Anlage, übergebene Unterlagen, Prüfergebnisse, Mängelliste
        mit freien Zeilen, Vorbehalte, Abnahmeentscheid und Unterschriften."""
        if protocol is None:
            protocol = self.create_acceptance_protocol()
        project = self.project
        client = getattr(project, "client_profile", None)
        company = self._company_profile

        pdf = self._make_pdf("Abnahmeprotokoll")
        pdf.add_heading("Abnahmeprotokoll KNX-Installation", level=1)

        # Projektdaten
        pdf.add_heading("Projektdaten", level=2)
        integrator = protocol.integrator_name
        if protocol.integrator_role:
            integrator += f" ({protocol.integrator_role})"
        contact = " | ".join(p for p in ((company.phone, company.email) if company else ()) if p)
        rows = [
            ["Projekt", project.name],
            ["Projektnummer", project.project_number or "–"],
            ["Objekt", (client.object_address if client else "") or "–"],
            ["Bauherr / Auftraggeber", protocol.client_name or "–"],
            ["Integrator", (integrator + (f"\n{contact}" if contact else "")) or "–"],
            ["Datum der Abnahme", protocol.date or datetime.now().strftime("%d.%m.%Y")],
        ]
        pdf.add_table(["Angabe", ""], rows, col_widths=[0.28, 0.72])

        # Umfang der Anlage
        pdf.add_heading("Umfang der Anlage", level=2)
        pdf.add_table(["", "Anzahl", "", "Anzahl"], self._plant_scope(),
                      col_widths=[0.32, 0.18, 0.32, 0.18],
                      align=["left", "right", "left", "right"])
        pdf.add_heading("Übergebene Unterlagen", level=3)
        docs = ["Bedienungsanleitung", "Inbetriebnahme-Checkliste",
                "Revisionsunterlagen (Berichte, Gruppenadressen)", "ETS-Projektdatei"]
        if project.knx_secure and project.knx_secure.enabled:
            docs.append("KNX-Secure-Archiv (vertraulich, separat)")
        pdf.add_paragraph("\n".join(f"[ ]  {d}" for d in docs))

        # Prüfergebnisse aus der Inbetriebnahme
        pdf.add_heading("Prüfergebnisse der Inbetriebnahme", level=2)
        items = [i for cl in project.checklists for i in cl.items]
        if items:
            ok = sum(1 for i in items if i.result == "OK")
            defects = sum(1 for i in items if i.result == "Mangel")
            na = sum(1 for i in items if i.result == "n/a")
            pdf.add_table(["Prüfpunkte gesamt", "in Ordnung", "Mängel", "nicht anwendbar", "offen"],
                          [[str(len(items)), str(ok), str(defects), str(na),
                            str(len(items) - ok - defects - na)]],
                          align=["right"] * 5)
        else:
            pdf.add_paragraph("Die Prüfergebnisse sind in der beiliegenden "
                              "Inbetriebnahme-Checkliste festgehalten.")

        # Mängelliste: erfasste Mängel + freie Zeilen
        pdf.add_heading("Mängelliste", level=2)
        rows = []
        for d in protocol.defects:
            rows.append([str(d.number), d.room_name, d.description, d.due_date,
                         "ja" if d.status == "behoben" else ""])
        for cl in project.checklists:
            for i in cl.items:
                if i.result == "Mangel":
                    where = " · ".join(p for p in (cl.room_name, i.be_number, i.check_type) if p)
                    rows.append([str(len(rows) + 1), where, i.notes or i.description, "", ""])
        for _ in range(max(6, 10 - len(rows)) if len(rows) < 10 else 2):
            rows.append([str(len(rows) + 1), "", "", "", ""])
        pdf.add_table(["Nr.", "Raum / Gerät", "Mangel", "Frist", "erledigt"], rows,
                      col_widths=[0.06, 0.26, 0.44, 0.12, 0.12])

        # Vorbehalte und Entscheid
        pdf.add_conditional_break(min_height=300)
        pdf.add_heading("Vorbehalte / Bemerkungen", level=2)
        if protocol.notes:
            pdf.add_paragraph(protocol.notes)
        # Schreiblinien zum Ausfüllen von Hand
        pdf.add_paragraph("\n".join(["", "_" * 118, "", "_" * 118, "", "_" * 118]))
        pdf.add_heading("Abnahmeentscheid", level=2)
        chosen = protocol.result
        options = ["Abnahme erfolgt", "Abnahme unter Vorbehalt (Mängel gemäss Liste)",
                   "Abnahme verweigert"]
        pdf.add_paragraph("\n".join(
            f"[{'x' if chosen and o.startswith(chosen) else ' '}]  {o}" for o in options))

        # Unterschriften
        pdf.add_heading("Unterschriften", level=2)
        pdf.add_table(
            ["", "Integrator", "Bauherr / Auftraggeber"],
            [["Name", protocol.integrator_name, protocol.client_name],
             ["Ort, Datum", "", ""],
             # geschützte Leerzeichen: Platz für die Unterschrift
             ["Unterschrift", " \n \n \n ", " \n \n \n "]],
            col_widths=[0.18, 0.41, 0.41])

        pdf.save(filepath)
        logger.info(f"Abnahmeprotokoll exportiert: {filepath}")

    # -- Bedienungsanleitung (FA-2001) --

    def generate_user_manual(self, filepath: str, language: str | None = None,
                             custom_intro: str = "",
                             include_scenes: bool = True):
        """Erzeugt die Bedienungsanleitung für den Bauherrn (FA-2001 bis
        FA-2006), siehe services/user_manual.py. language: "de", "fr", "it"
        oder "en"; ohne Angabe die Sprache aus «Anleitung anpassen».
        include_scenes bleibt für bestehende Aufrufe erhalten."""
        from .user_manual import UserManualBuilder
        pdf = self._make_pdf("Bedienungsanleitung")
        UserManualBuilder(self.project, self._company_profile,
                          language=language).build(pdf, custom_intro)
        pdf.save(filepath)
        logger.info(f"Bedienungsanleitung erstellt: {filepath}")

    # -- Revisionspaket (FA-2100) --

    def generate_revision_package(self, output_dir: str, revision: str = "",
                                   language: str | None = None, parts=None,
                                   revision_date: str = "", note: str = ""):
        """Erzeugt ein komplettes Revisionspaket (FA-2101 bis FA-2106).

        parts: Schlüssel aus revision_package.REVISION_PARTS (FA-2105);
            None = alle. Bestandteile ohne Daten entfallen automatisch.
        revision, revision_date, note: Revisionsstand (FA-2106), erscheinen
            im Inhaltsverzeichnis.
        Mit `revision` wird der Stand in project.revisions festgehalten."""
        from .revision_package import (
            ALL_PARTS, accepted_quotes, collect_datasheets, copy_datasheets,
        )
        os.makedirs(output_dir, exist_ok=True)
        wanted = ALL_PARTS if parts is None else frozenset(parts)

        project_name = self.project.name or "Projekt"
        prefix = project_name.replace(" ", "_")
        if revision:
            prefix = f"{prefix}_Rev{revision}"

        generated_files = []

        def target(suffix: str) -> str:
            return os.path.join(output_dir, f"{prefix}_{suffix}")

        from .report_service import ReportService
        report_svc = ReportService(self.project, company_profile=self._company_profile)

        # 1. Projektzusammenfassung
        if "zusammenfassung" in wanted:
            path = target("Zusammenfassung.pdf")
            report_svc.generate_project_summary(path)
            generated_files.append(("Projektzusammenfassung", path))

        # 2. GA-Übersicht
        if "ga_uebersicht" in wanted:
            path = target("GA_Uebersicht.pdf")
            report_svc.generate_ga_report(path)
            generated_files.append(("Gruppenadress-Übersicht", path))

        # 3. Topologie-Bericht
        if "topologie" in wanted:
            path = target("Topologie.pdf")
            report_svc.generate_topology_report(path)
            generated_files.append(("Topologie-Bericht", path))

        # 4. Bedienelemente-Bericht
        if "bedienelemente" in wanted:
            path = target("Bedienelemente.pdf")
            report_svc.generate_bedienelemente_report(path)
            generated_files.append(("Bedienelemente und Sensoren", path))

        # 4b. Aktoren und Gateways
        if "aktoren" in wanted:
            path = target("Aktoren_Gateways.pdf")
            report_svc.generate_aktoren_gateway_report(path)
            generated_files.append(("Aktoren und Gateways", path))

        # 4c. Räume nach Gewerken
        if "raeume" in wanted:
            path = target("Raeume_Gewerke.pdf")
            report_svc.generate_room_gewerk_report(path)
            generated_files.append(("Räume nach Gewerken", path))

        # 4d. Verknüpfungsmatrix / Belegungsplan (FA-2505) -- nur wenn Daten vorhanden
        if "belegungsplan" in wanted:
            from .belegungsplan_service import BelegungsplanService
            from .sensor_service import project_for_export
            belegungsplan = BelegungsplanService().generate(project_for_export(self.project))
            if belegungsplan.sensor_rows or belegungsplan.actor_rows:
                from .belegungsplan_export_service import BelegungsplanExportService
                path = target("Verknuepfungsmatrix.pdf")
                BelegungsplanExportService().export_pdf(
                    belegungsplan, path, self._company_profile, self.project.project_info
                )
                generated_files.append(("Belegungsplan (Verknüpfungsmatrix)", path))

        # 4e. Szenenreport (FA-1811) -- nur wenn Szenen definiert sind
        if "szenen" in wanted and any(s.name for s in self.project.scenes):
            path = target("Szenenreport.pdf")
            report_svc.generate_szenen_report(path)
            generated_files.append(("Szenenreport", path))

        # 5. Validierungsbericht
        if "validierung" in wanted:
            path = target("Validierung.pdf")
            report_svc.generate_validation_report(path)
            generated_files.append(("Validierungsbericht Gruppenadressen", path))

        # 6. Inbetriebnahme-Checklisten
        if "checklisten" in wanted:
            checklists = self.create_checklists()
            path = target("Checklisten.pdf")
            self.export_checklists_pdf(path, checklists)
            generated_files.append(("Inbetriebnahme-Checklisten", path))

        # 6b. Abnahmeprotokoll (Formular zum Unterschreiben)
        if "abnahme" in wanted:
            path = target("Abnahmeprotokoll.pdf")
            self.export_acceptance_protocol(path, self.project.acceptance_protocol)
            generated_files.append(("Abnahmeprotokoll", path))

        # 7. Bedienungsanleitung
        if "anleitung" in wanted:
            path = target("Bedienungsanleitung.pdf")
            self.generate_user_manual(path, language=language)
            generated_files.append(("Bedienungsanleitung", path))

        # 7b. Materialliste (FA-2102 Nr. 12, FA-2308) -- nur wenn erfasst
        if "materialliste" in wanted and self.project.material_list.entries:
            from .material_list_export_service import MaterialListExportService
            path = target("Materialliste.xlsx")
            MaterialListExportService().export_xlsx(
                self.project.material_list, project_name, path)
            generated_files.append(("Materialliste", path))

        # 8. CSV-Export
        if "ga_csv" in wanted:
            from .csv_export_service import CsvExportService
            path = target("GA_Export.csv")
            CsvExportService().export_csv(
                self.project.group_addresses, path, overwrite=True
            )
            generated_files.append(("GA-Export (CSV)", path))

        # 9. DALI-Geräteliste (FA-2805) – nur wenn DALI-Konfigurationen vorhanden
        if "dali" in wanted and any(gw.devices for gw in self.project.dali_configs.values()):
            path = target("DALI_Geraete.pdf")
            self.generate_dali_device_list(path)
            generated_files.append(("DALI-Gerätekonfiguration", path))

        # 10. Zeitprogramme (FA-3307) – nur wenn Zeitprogramme vorhanden
        if "zeitprogramme" in wanted and self.project.time_programs:
            path = target("Zeitsteuerungsplan.pdf")
            self.generate_time_programs_doc(path)
            generated_files.append(("Zeitsteuerungsplan", path))

        # 10b. KNX Secure Archivbericht (FA-2706) – nur wenn aktiviert
        if "secure" in wanted and self.project.knx_secure.enabled:
            path = target("KNX_Secure.pdf")
            self.generate_knx_secure_report(path)
            generated_files.append(("KNX Secure Archivbericht", path))

        # 10c. Kundenofferte (FA-2102 Nr. 13) – nur akzeptierte
        if "offerte" in wanted:
            from .quote_letter_service import write_quote_letter_pdf
            quotes = accepted_quotes(self.project)
            for i, quote in enumerate(quotes, 1):
                suffix = f"_{i}" if len(quotes) > 1 else ""
                path = target(f"Kundenofferte{suffix}.pdf")
                write_quote_letter_pdf(self.project, quote, project_name, path)
                generated_files.append(("Kundenofferte (akzeptiert)", path))

        # 10d. Produktdatenblätter (FA-1204) als Anhang
        datasheet_rows = []
        if "datenblaetter" in wanted:
            datasheets = collect_datasheets(self.project)
            copied, missing = copy_datasheets(
                datasheets, output_dir, self.project.folder_path or "")
            for ds, path in copied:
                datasheet_rows.append([", ".join(ds.products),
                                       os.path.relpath(path, output_dir)])
            for ds in datasheets:
                if ds.is_url:
                    datasheet_rows.append([", ".join(ds.products), ds.path])
            for ds in missing:
                datasheet_rows.append([", ".join(ds.products),
                                       f"nicht gefunden: {ds.path}"])

        # 11. Inhaltsverzeichnis (FA-2103) mit offenen Punkten (FA-2104)
        from .revision_check import check_revision_completeness
        self.create_revision_index(
            output_dir, revision, generated_files,
            findings=check_revision_completeness(self.project, self._company_profile),
            revision_date=revision_date, note=note, datasheet_rows=datasheet_rows,
        )

        # Revisionsstand festhalten (FA-2106); dieselbe Bezeichnung ersetzt
        # einen früheren Stand, z.B. beim erneuten Erstellen von Rev. B
        if revision:
            from ..models.documentation import RevisionRecord
            parts_done = list(dict.fromkeys(name for name, _ in generated_files))
            if datasheet_rows:
                parts_done.append("Produktdatenblätter")
            self.project.revisions = [r for r in self.project.revisions
                                      if r.number != revision]
            self.project.revisions.append(RevisionRecord(
                revision, revision_date or datetime.now().strftime("%d.%m.%Y"),
                note, parts_done))

        logger.info(f"Revisionspaket erstellt in: {output_dir}")
        return output_dir

    def generate_knx_secure_report(self, filepath: str, include_secrets: bool = False):
        """Erzeugt den KNX-Secure-Archivbericht (FA-2706).

        Dokumentiert den Erfassungsstatus des Zertifikats- und Zugangsdaten-
        Archivs (FDSK je Gerät, ETS6-Projektpasswort) für die Übergabe.

        include_secrets=False (Standard, z.B. im automatischen Revisionspaket):
        die eigentlichen Geheimwerte (FDSK, ETS6-Projektpasswort) werden
        bewusst NICHT abgedruckt -- nur ob sie erfasst sind.

        include_secrets=True (nur auf expliziten Nutzerwunsch, siehe
        ReportsDialog._gen_knx_secure_confidential -- dort wird die Kenntnis
        des ETS6-Projektpassworts vor Erstellung geprüft): druckt die echten
        FDSK- und Passwort-Werte ab, als Vertraulichkeits-Archivkopie für den
        Fall, dass das Master-Passwort des Archivs verloren geht. Der Bericht
        ist dann selbst ein Geheimnis und muss entsprechend vertraulich
        behandelt werden.
        """
        from ..services.knx_secure_service import KnxSecureService, sorted_device_infos
        from ..models.knx_secure import SECURE_MODE_LABELS

        cfg = self.project.knx_secure
        svc = KnxSecureService()

        title = "KNX Secure – Archivbericht (VERTRAULICH)" if include_secrets else "KNX Secure – Archivbericht"
        pdf = self._make_pdf(title)
        pdf.add_heading("KNX Secure – Zertifikats- und Zugangsdaten-Archiv", level=1)
        if include_secrets:
            pdf.add_heading("⚠ VERTRAULICH -- enthält echte FDSK- und Passwort-Werte", level=2)
        pdf.add_paragraph(f"Projekt: {self.project.name}")
        pdf.add_paragraph(f"Erstellt: {datetime.now().strftime('%d.%m.%Y %H:%M')}")

        if not cfg.enabled:
            pdf.add_paragraph("KNX Secure ist für dieses Projekt nicht aktiviert.")
            pdf.save(filepath)
            return

        pdf.add_paragraph(
            "Hinweis: KNiX Arranger generiert und verwaltet keine KNX-Secure-"
            "Laufzeitschlüssel -- diese erzeugt ausschliesslich ETS6 aus dem "
            "geräteindividuellen Zertifikat (FDSK). " + (
                "Dieser Bericht enthält die tatsächlich hinterlegten Geheimwerte "
                "als Archivkopie -- vertraulich behandeln, nicht ungeschützt "
                "weitergeben oder ablegen."
                if include_secrets else
                "Dieser Bericht dokumentiert nur den Erfassungsstatus des "
                "Archivs, keine Geheimwerte."
            )
        )

        if cfg.is_locked:
            pdf.add_paragraph(
                "⚠ Das Archiv war beim Erstellen dieses Berichts gesperrt "
                "(Master-Passwort nicht eingegeben). Angaben zu FDSK- und "
                "Passwort-Erfassung unten können daher unvollständig sein."
            )

        summary = svc.get_summary(cfg, self.project)
        mode_label = SECURE_MODE_LABELS.get(cfg.secure_mode, cfg.secure_mode)
        pdf.add_separator()
        pdf.add_heading("Projektübersicht", level=2)
        pdf.add_paragraph(f"Sicherheitsmodus: {mode_label}")
        if include_secrets and summary['ets_password_set']:
            pdf.add_paragraph(f"ETS6-Projektpasswort: {cfg.ets_project_password}")
        else:
            pdf.add_paragraph(
                "ETS6-Projektpasswort hinterlegt: "
                f"{'Ja' if summary['ets_password_set'] else 'Nein'}"
            )
        if cfg.ets_password_note:
            pdf.add_paragraph(f"Notiz / Wiederherstellungsplan: {cfg.ets_password_note}")
        pdf.add_paragraph(
            f"KNX Secure-fähige Geräte: {summary['secure_devices']} / {summary['device_infos_count']}"
        )
        pdf.add_paragraph(
            f"FDSK-Zertifikate erfasst: {summary['fdsk_entered']} / {summary['device_infos_count']}"
        )
        pdf.add_paragraph(
            f"Gruppenadressen mit Security 'Ein': {summary['ga_security_on']}"
        )

        pdf.add_separator()
        pdf.add_heading("Geräte", level=2)
        infos = sorted_device_infos(cfg)
        if not infos:
            pdf.add_paragraph("Keine Geräte im Archiv erfasst.")
        else:
            # Adresse, Linie und Gerätename aus der aktuellen Topologie: das
            # Archiv kann nach einem Neuaufbau veraltete Linien-IDs und den
            # englischen Produktnamen aus der ETS tragen
            by_id = {d.id: (f"{a.area_number}.{ln.line_number}", d)
                     for a in self.project.topology.areas for ln in a.lines
                     for d in ln.devices}
            id_of = {id(v): k for k, v in cfg.device_infos.items()}
            headers = ["Gerät", "Adresse", "Linie", "KNX Secure",
                       "FDSK" if include_secrets else "FDSK erfasst"]
            rows = []
            for info in infos:
                line_label, device = by_id.get(id_of.get(id(info), ""), ("", None))
                if device is not None:
                    name = device.product_name or device.product or info.device_name
                    address = device.physical_address or info.physical_address
                else:
                    name = f"{info.device_name} (nicht mehr in der Topologie)"
                    address = info.physical_address
                rows.append([
                    name, address, line_label,
                    "Ja" if info.secure_supported else "Nein",
                    (info.fdsk or "–") if include_secrets else ("Ja" if info.fdsk else "Nein"),
                ])
            rows.sort(key=lambda r: physical_address_key(r[1]))
            pdf.add_table(headers, rows, col_widths=[0.46, 0.12, 0.10, 0.14, 0.18])

        warnings = svc.check_mixed_lines(cfg, self.project)
        pdf.add_separator()
        pdf.add_heading("Mischlinien-Warnungen", level=2)
        if not warnings:
            pdf.add_paragraph("Keine Mischlinien-Warnungen.")
        else:
            for w in warnings:
                pdf.add_paragraph(f"• {w.message}")

        pdf.save(filepath)
        logger.info(f"KNX-Secure-Archivbericht erstellt: {filepath}")

    def generate_dali_device_list(self, filepath: str):
        """Erzeugt die DALI-Geräteliste für das Revisionspaket (FA-2805)."""
        from .dali_service import DaliService
        svc = DaliService()
        # Gateways nach physikalischer Adresse (dali_configs ist nach
        # Geraete-ID indiziert)
        gateway_address = {
            d.id: d.physical_address
            for area in self.project.topology.areas for line in area.lines for d in line.devices
        }
        existing_gas = {g.address: g for g in self.project.group_addresses.all_addresses()}
        pdf = self._make_pdf("DALI-Gerätekonfiguration")
        pdf.add_heading("DALI-Gerätekonfiguration", level=1)
        pdf.add_paragraph(f"Projekt: {self.project.name}")
        pdf.add_paragraph(f"Erstellt: {datetime.now().strftime('%d.%m.%Y %H:%M')}")

        if not self.project.dali_configs:
            pdf.add_paragraph("Keine DALI-Konfigurationen vorhanden.")
        else:
            for gw_id, gw in sorted(
                self.project.dali_configs.items(),
                key=lambda item: physical_address_key(gateway_address.get(item[0], "")),
            ):
                pdf.add_separator()
                address = gateway_address.get(gw_id, "")
                pdf.add_heading(" · ".join(p for p in (address, gw.name or "DALI-Gateway") if p),
                                level=2)
                # Nur Gruppenadressen, die es im Projekt gibt (Vorgaben des
                # Assistenten ohne GA, z.B. in importierten Projekten, entfallen)
                for label, addr in (("Schalten (Broadcast)", gw.ga_switch_broadcast),
                                    ("Dimmen (Broadcast)", gw.ga_dim_broadcast),
                                    ("Szene", gw.ga_scene)):
                    ga = existing_gas.get(addr) if addr else None
                    if ga is not None:
                        designation = " ".join((ga.designation or "").split())
                        pdf.add_paragraph(f"GA {label}: {addr}  {designation}")

                # Gruppen mit ihren GAs (für die Parametrierung des Gateways)
                if gw.groups:
                    pdf.add_heading("Gruppen", level=3)
                    pdf.add_table(
                        ["Gr.", "Name", "Schalten", "Dimmen", "Wert", "Status", "Szene",
                         "Störung"],
                        [[str(g.number), g.name, g.ga_switch, g.ga_dim, g.ga_value,
                          g.ga_status, g.ga_scene, g.ga_fault]
                         for g in sorted(gw.groups, key=lambda g: g.number)],
                        col_widths=[0.06, 0.26, 0.11, 0.11, 0.11, 0.11, 0.12, 0.12],
                    )

                rows = svc.generate_device_list(gw, self.project)
                if rows:
                    headers = ["Adr.", "Name", "EVG-Typ", "Raum", "Gruppen", "Notlicht", "Modus"]
                    table_rows = [
                        [
                            str(r["short_address"]), r["name"], r["evg_type"],
                            r["room"], r["groups"],
                            "Ja" if r["is_emergency"] else "",
                            r["emergency_mode"] if r["is_emergency"] else "",
                        ]
                        for r in rows
                    ]
                    pdf.add_table(headers, table_rows)
                else:
                    pdf.add_paragraph("Keine EVGs in KNiX erfasst (Konfiguration im Gateway "
                                      "bzw. in der Inbetriebnahme-Software des Herstellers).")

        pdf.save(filepath)
        logger.info(f"DALI-Geräteliste erstellt: {filepath}")

    def generate_time_programs_doc(self, filepath: str):
        """Zeitsteuerungsplan: alle Wochenprogramme und Schaltzeitpunkte
        (FA-3307b/c), auch im Revisionspaket."""
        pdf = self._make_pdf("Zeitsteuerungsplan")
        pdf.add_heading("Zeitsteuerungsplan", level=1)
        pdf.add_paragraph(f"Projekt: {self.project.name}")
        pdf.add_paragraph(f"Erstellt: {datetime.now().strftime('%d.%m.%Y %H:%M')}")

        loc = self.project.location
        pdf.add_paragraph(
            f"Standort: {loc.latitude:.4f}°N / {loc.longitude:.4f}°O "
            f"({loc.country}, {loc.region})"
        )

        if not self.project.time_programs:
            pdf.add_paragraph("Keine Zeitprogramme konfiguriert.")
        else:
            for tp in self.project.time_programs:
                pdf.add_separator()
                status = "aktiv" if tp.active else "inaktiv"
                pdf.add_heading(f"{tp.name} ({status})", level=2)
                pdf.add_paragraph(
                    f"Schaltzeitpunkte gesamt: {tp.switch_point_count}  |  "
                    f"Tagesprofile: {len(tp.day_profiles)}"
                )
                for i, dp in enumerate(tp.day_profiles):
                    days = ", ".join(dp.weekdays) or "—"
                    pdf.add_heading(f"Tagesprofil {i+1}: {days}", level=3)
                    if not dp.switch_points:
                        pdf.add_paragraph("Keine Schaltzeitpunkte.")
                        continue
                    headers = ["Zeit", "Art", "Wert", "GA", "Priorität", "Zeitraum"]
                    rows = []
                    for sp in sorted(dp.switch_points,
                                     key=lambda s: s.fixed_time if s.time_type == "FIXED" else ""):
                        ga_label = ""
                        if sp.target_ga_id:
                            ga = next(
                                (g for g in self.project.group_addresses.all_addresses()
                                 if g.id == sp.target_ga_id), None
                            )
                            ga_label = (
                                f"{ga.main_group}/{ga.middle_group}/{ga.sub_group} "
                                f"{ga.designation}"
                            ) if ga else sp.target_ga_id
                        rows.append([
                            sp.display_time,
                            "Astro" if sp.time_type == "ASTRO" else "fest",
                            sp.action_value, ga_label, sp.priority,
                            _period_text(sp),
                        ])
                    pdf.add_table(headers, rows,
                                  col_widths=[0.16, 0.08, 0.08, 0.40, 0.10, 0.18])

        pdf.save(filepath)
        logger.info(f"Zeitprogramme-Dokumentation erstellt: {filepath}")

    def create_revision_index(self, output_dir: str, revision: str,
                              generated_files: list[tuple[str, str]],
                              findings: list | None = None,
                              revision_date: str = "", note: str = "",
                              datasheet_rows: list | None = None) -> str:
        """Erzeugt ein Inhaltsverzeichnis für das Revisionspaket (FA-2103)."""
        project_name = self.project.name or "Projekt"
        prefix = project_name.replace(" ", "_")
        if revision:
            prefix = f"{prefix}_Rev{revision}"

        index_path = os.path.join(output_dir, f"{prefix}_Inhaltsverzeichnis.pdf")

        pdf = self._make_pdf("Revisionspaket - Inhaltsverzeichnis")

        pdf.add_heading("Revisionspaket - Inhaltsverzeichnis", level=1)
        pdf.add_paragraph(f"Projekt: {self.project.name}")
        if revision:
            date = revision_date or datetime.now().strftime("%d.%m.%Y")
            pdf.add_paragraph(f"Revision: {revision} vom {date}")
        if note:
            pdf.add_paragraph(f"Anlass: {note}")
        pdf.add_paragraph(f"Erstellt: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
        pdf.add_separator()

        headers = ["Nr.", "Dokument", "Seiten", "Datei"]
        rows = []
        for i, (doc_name, filepath) in enumerate(generated_files, 1):
            pages = ""
            if filepath.lower().endswith(".pdf"):
                try:
                    import fitz
                    with fitz.open(filepath) as doc:
                        pages = str(len(doc))
                except Exception:
                    pages = ""
            # Projektpräfix weglassen: der Dateiname bricht sonst mitten im Wort um
            name = os.path.basename(filepath)
            short = name[len(prefix) + 1:] if name.startswith(prefix + "_") else name
            rows.append([str(i), doc_name, pages, short])
        pdf.add_table(headers, rows, col_widths=[0.07, 0.43, 0.10, 0.40],
                      align=["right", "left", "right", "left"])
        pdf.add_paragraph(f"Alle Dateinamen beginnen mit «{prefix}_».")

        # Produktdatenblätter (FA-1204)
        if datasheet_rows:
            pdf.add_heading("Produktdatenblätter", level=2)
            pdf.add_table(["Produkt", "Datei / Link"], datasheet_rows,
                          col_widths=[0.40, 0.60])

        # Frühere Revisionen (FA-2106)
        earlier = [r for r in self.project.revisions if r.number != revision]
        if earlier:
            pdf.add_heading("Frühere Revisionen", level=2)
            pdf.add_table(["Revision", "Datum", "Anlass"],
                          [[r.number, r.date, r.note] for r in earlier],
                          col_widths=[0.15, 0.20, 0.65])

        # Vollständigkeit (FA-2104)
        pdf.add_heading("Offene Punkte", level=2)
        if findings:
            pdf.add_table(["Bestandteil", "Offen"],
                          [[f.part, f.message] for f in findings],
                          col_widths=[0.28, 0.72])
        else:
            pdf.add_paragraph("Die Revisionsunterlagen sind vollständig.")

        pdf.save(index_path)
        logger.info(f"Inhaltsverzeichnis erstellt: {index_path}")
        return index_path
