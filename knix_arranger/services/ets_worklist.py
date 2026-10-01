"""
ETS-Arbeitsliste (FA-618)

Bei importierten Projekten findet KNiX Fehler, behoben werden sie in der
ETS (KNiX kann nichts in die ETS zurückschreiben). Die Arbeitsliste sagt dem
Integrator knapp und in ETS-Begriffen, was er wo in der ETS tun muss:

1. Muss behoben werden: wirkt auf den Bus (z.B. zweite sendende GA an einer
   Taste, eine in KNiX getrennte GA, die in der ETS noch verbunden ist).
2. Prüfen: Warnungen der Validierung (z.B. fehlender Datenpunkttyp).
3. Unterlagen nachführen: in KNiX korrigierte Angaben (Gewerk, Raum), die
   nur die Dokumentation betreffen -- kein Programmieren nötig.

Der Stand wird im Projekt gemerkt (project.ets_worklist). Nach jedem
Re-Import gleicht sync_worklist ab: was die Validierung nicht mehr meldet,
ist in der ETS erledigt; was neu dazukommt, ist als "neu" markiert.
"""
from __future__ import annotations
import logging
from dataclasses import dataclass, field
from datetime import date

logger = logging.getLogger("knix_arranger.ets_worklist")

PRIO_MUST, PRIO_CHECK, PRIO_DOCS = 1, 2, 3
PRIORITIES = {
    PRIO_MUST: ("Muss behoben werden", "Wirkt auf den Bus – nach der Korrektur programmieren."),
    PRIO_CHECK: ("Prüfen", "Verdächtig, aber nicht eindeutig falsch."),
    PRIO_DOCS: ("Unterlagen nachführen", "Nur Dokumentation – kein Programmieren nötig."),
}

STATUS_NEW, STATUS_OPEN, STATUS_DONE = "neu", "offen", "erledigt"

# Kurz: wer die ETS bedient, kennt die Handgriffe (Grundkurs). Begriffe wie
# in der ETS ("trennen", nicht "Verbindung löschen").
HOW_TO_USE = ("In der ETS korrigieren, betroffene Geräte programmieren, Projekt exportieren "
              "und in KNiX erneut importieren – KNiX hakt Erledigtes selbst ab.")


@dataclass
class WorkItem:
    key: str
    priority: int
    rule_id: str
    where: str          # Ort in der ETS: "1.1.51 KO 6 «Taste 2, links»" bzw. "GA 0/0/101"
    task: str           # was zu tun ist, in ETS-Begriffen
    detail: str = ""    # Bezeichnung
    status: str = STATUS_OPEN
    first_seen: str = ""
    done_on: str = ""
    sort: tuple = field(default_factory=tuple)


@dataclass
class Worklist:
    items: list[WorkItem]           # offen und neu
    done: list[WorkItem]            # erledigt
    last_import: str = ""

    def by_priority(self) -> dict[int, list[WorkItem]]:
        groups: dict[int, list[WorkItem]] = {p: [] for p in PRIORITIES}
        for item in self.items:
            groups[item.priority].append(item)
        return groups

    @property
    def new_count(self) -> int:
        return sum(1 for i in self.items if i.status == STATUS_NEW)


# ── Aufgaben aus der Validierung ──────────────────────────────────────────

def _addr_sort(text: str) -> tuple:
    import re
    return tuple(int(n) for n in re.findall(r"\d+", text)[:4])


def _gewerk_rename(designation: str, ets_code: str, knix_code: str) -> str:
    if designation.startswith(f"{ets_code}."):
        return f"{knix_code}." + designation[len(ets_code) + 1:]
    return ""


def _items_from_issue(issue, project) -> list[WorkItem]:
    rule = issue.rule_id
    details = issue.details or {}
    addr = issue.address
    designation = (issue.designation or "").strip()

    if rule == "FA-614":
        ko = details.get("ko", "")              # "1.1.51 KO 6 «Taste 2, links»"
        if issue.level == "error":
            return [WorkItem(
                f"FA-614|{ko}|{addr}", PRIO_MUST, rule, ko,
                f"{addr} trennen (zweite sendende GA).", designation,
                sort=(PRIO_MUST, 0, _addr_sort(ko), _addr_sort(addr)))]
        # Warnung: kein Sender gefunden -- die Taste hört die GA nur mit
        return [WorkItem(
            f"FA-614|{ko}|{addr}", PRIO_CHECK, rule, ko,
            f"{addr} hört mit, kein Sender gefunden – gewollt? Sonst trennen.", designation,
            sort=(PRIO_CHECK, 0, _addr_sort(ko), _addr_sort(addr)))]

    if rule == "FA-616":
        what = details.get("field", "")
        ets, knix = details.get("ets", ""), details.get("knix", "")
        if what == "Verknüpfung":
            # designation "1.1.51 KO 6 · S1 Wohnen_ea"
            ref, _, ga_name = designation.partition(" · ")
            pa, _, co_label = ref.partition(" ")
            from .multi_ga_check import find_device
            device = find_device(project, pa)
            co_no = co_label.replace("KO", "").strip()
            co = None if device is None else next(
                (c for c in device.communication_objects if str(c.object_number) == co_no), None)
            if co is not None and (co.name or co.object_function):
                co_label = f"{co_label} «{co.name or co.object_function}»"
            return [WorkItem(
                f"FA-616|Verknüpfung|{pa}|{co_no}|{addr}", PRIO_MUST, rule,
                f"{pa} {co_label}", f"{addr} trennen (in KNiX bereits getrennt).", ga_name,
                sort=(PRIO_MUST, 1, _addr_sort(pa), _addr_sort(addr)))]
        if what == "Raum":
            return [WorkItem(
                f"FA-616|Raum|{addr}", PRIO_DOCS, rule, f"Gerät {addr}",
                f"In Raum «{knix}» verschieben.", f"Heute: «{ets}»",
                sort=(PRIO_DOCS, 1, _addr_sort(addr)))]
        new_name = _gewerk_rename(designation, ets, knix)
        return [WorkItem(
            f"FA-616|{what}|{addr}", PRIO_DOCS, rule, f"GA {addr}",
            f"Umbenennen in «{new_name}»." if new_name else f"Gewerk-Kürzel {ets} → {knix}.",
            f"Heute: {designation}" if designation else "",
            sort=(PRIO_DOCS, 0, _addr_sort(addr)))]

    if rule in ("FA-603", "FA-604"):
        proposal = details.get("dpt") or details.get("expected") or ""
        task = (("Datenpunkttyp setzen." if not proposal else
                 f"Datenpunkttyp {proposal} setzen (Vorschlag).") if rule == "FA-604" else
                f"Datenpunkttyp prüfen: {details.get('actual', '')}, erwartet {proposal}.")
        return [WorkItem(
            f"{rule}|{addr}", PRIO_CHECK, rule, f"GA {addr}", task, designation,
            sort=(PRIO_CHECK, 1 if rule == "FA-604" else 0, _addr_sort(addr)))]

    # Übrige Regeln: Meldung der Validierung
    prio = PRIO_MUST if issue.level == "error" else PRIO_CHECK
    where = f"GA {addr}" if addr and "/" in addr else (addr or "Projekt")
    return [WorkItem(
        f"{rule}|{addr}|{issue.message}", prio, rule, where, issue.message, designation,
        sort=(prio, 2, rule, _addr_sort(addr or "")))]


def current_items(project) -> list[WorkItem]:
    """Aufgaben aus dem aktuellen Projektstand: Fehler und Warnungen der
    Validierung sowie Abweichungen zur ETS. Hinweise (z.B. Namensschema bei
    importierten Projekten) gehören nicht dazu."""
    from .validation_engine import ValidationEngine
    issues = ValidationEngine(project.gewerk_catalog).validate(
        project.group_addresses, project=project)
    items: dict[str, WorkItem] = {}
    for issue in issues:
        if issue.level == "info" and issue.rule_id != "FA-616":
            continue
        for item in _items_from_issue(issue, project):
            items.setdefault(item.key, item)
    return sorted(items.values(), key=lambda i: i.sort)


# ── Stand merken und abgleichen ───────────────────────────────────────────

def _state(project) -> dict:
    state = project.ets_worklist
    state.setdefault("round", 0)
    state.setdefault("last_import", "")
    state.setdefault("items", {})
    return state


def build_worklist(project) -> Worklist:
    """Arbeitsliste zum aktuellen Stand, verglichen mit dem letzten Import
    (verändert das Projekt nicht)."""
    state = _state(project)
    stored = state["items"]
    current = current_items(project)
    current_keys = {i.key for i in current}
    for item in current:
        entry = stored.get(item.key)
        if entry and not entry.get("done_on"):
            item.first_seen = entry.get("first_seen", "")
            item.status = (STATUS_NEW if state["round"] > 1
                           and entry.get("round") == state["round"] else STATUS_OPEN)
        else:
            item.status = STATUS_NEW if state["round"] else STATUS_OPEN
    done = []
    for key, entry in stored.items():
        if key in current_keys:
            continue
        done.append(WorkItem(
            key, entry.get("priority", PRIO_CHECK), entry.get("rule_id", ""),
            entry.get("where", ""), entry.get("task", ""), entry.get("detail", ""),
            status=STATUS_DONE, first_seen=entry.get("first_seen", ""),
            done_on=entry.get("done_on", "") or "seit letztem Import",
            sort=(entry.get("priority", PRIO_CHECK), _addr_sort(entry.get("where", "")))))
    done.sort(key=lambda i: i.sort)
    return Worklist(current, done, state["last_import"])


@dataclass
class WorklistDiff:
    done: int = 0
    new: int = 0
    open: int = 0
    first: bool = False

    def text(self) -> str:
        return (f"{self.done} erledigt, {self.new} neu, {self.open} offen")


def sync_worklist(project, today: str | None = None) -> WorklistDiff:
    """Nach einem Import: Stand der Arbeitsliste fortschreiben. Was nicht mehr
    gemeldet wird, gilt als erledigt (mit Datum); Neues wird festgehalten."""
    today = today or date.today().strftime("%d.%m.%Y")
    state = _state(project)
    stored = state["items"]
    first = state["round"] == 0
    state["round"] += 1
    state["last_import"] = today
    current = current_items(project)
    diff = WorklistDiff(first=first)
    current_keys = set()
    for item in current:
        current_keys.add(item.key)
        entry = stored.get(item.key)
        if entry is None or entry.get("done_on"):
            stored[item.key] = {
                "priority": item.priority, "rule_id": item.rule_id, "where": item.where,
                "task": item.task, "detail": item.detail, "first_seen": today,
                "round": state["round"], "done_on": "",
            }
            diff.new += 1
        else:
            entry.update(priority=item.priority, where=item.where, task=item.task,
                         detail=item.detail)
            diff.open += 1
    for key, entry in stored.items():
        if key not in current_keys and not entry.get("done_on"):
            entry["done_on"] = today
            diff.done += 1
    if first:
        diff.open, diff.new = diff.new, 0
    return diff


def ensure_baseline(project) -> bool:
    """Vor dem ersten Erstellen der Liste den Ausgangsstand merken -- sonst
    könnte der nächste Re-Import nicht feststellen, was in der ETS erledigt
    wurde. Gibt True zurück, wenn das Projekt dadurch geändert wurde."""
    if _state(project)["round"]:
        return False
    sync_worklist(project)
    return True


# ── Ausgabe ───────────────────────────────────────────────────────────────

_PDF_HEADERS = ["Nr.", "Ort in der ETS", "Aufgabe", "Erledigt"]
_PDF_WIDTHS = [0.07, 0.24, 0.59, 0.10]


def _numbered(worklist: Worklist):
    n = 0
    for prio, items in worklist.by_priority().items():
        rows = []
        for item in items:
            n += 1
            rows.append((n, item))
        yield prio, rows


def _task_text(item: WorkItem) -> str:
    return item.task + (f"\n{item.detail}" if item.detail else "")


def export_pdf(project, filepath: str, company_profile=None) -> Worklist:
    from ..utils.pdf_generator import PdfGenerator
    worklist = build_worklist(project)
    pdf = PdfGenerator(title="ETS-Arbeitsliste", company_profile=company_profile,
                       project_info=project.project_info)
    pdf.add_heading("ETS-Arbeitsliste", level=1)
    since = f" | Verglichen mit Stand vom {worklist.last_import}" if worklist.last_import else ""
    pdf.add_paragraph(f"Projekt: {project.name} | Erstellt: "
                      f"{date.today().strftime('%d.%m.%Y')}{since}")
    pdf.add_paragraph(HOW_TO_USE)

    for prio, rows in _numbered(worklist):
        if not rows:
            continue
        title, explanation = PRIORITIES[prio]
        new = sum(1 for _n, i in rows if i.status == STATUS_NEW)
        pdf.add_conditional_break(min_height=120)
        pdf.add_heading(f"{prio}  {title}  ({len(rows)}" + (f", davon {new} neu)" if new else ")"),
                        level=2)
        pdf.add_paragraph(explanation)
        pdf.add_table(_PDF_HEADERS, [
            [("NEU " if item.status == STATUS_NEW else "") + str(n), item.where,
             _task_text(item), "[ ]"] for n, item in rows], col_widths=_PDF_WIDTHS)

    if worklist.done:
        pdf.add_conditional_break(min_height=120)
        pdf.add_heading(f"Erledigt  ({len(worklist.done)})", level=2)
        pdf.add_table(["Ort in der ETS", "Aufgabe", "Erledigt am"],
                      [[i.where, i.task, i.done_on] for i in worklist.done],
                      col_widths=[0.24, 0.60, 0.16])
    pdf.save(filepath)
    logger.info("ETS-Arbeitsliste PDF: %s", filepath)
    return worklist


XL_CHOICES = ("☐ offen", "✓ erledigt", "– später")


def export_excel(project, filepath: str) -> Worklist:
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import PatternFill
    from ..utils.excel_generator import ExcelGenerator, HAS_OPENPYXL
    if not HAS_OPENPYXL:
        raise ImportError("openpyxl wird für Excel-Export benötigt.")
    worklist = build_worklist(project)
    excel = ExcelGenerator(title="ETS-Arbeitsliste", project_name=project.name)
    excel._current_sheet.title = "Arbeitsliste"
    excel.set_column_widths([7, 9, 30, 70, 12, 34])
    excel.set_print_options(orientation="landscape")
    excel.add_header()
    excel.add_empty_row()
    excel.add_paragraph(HOW_TO_USE)
    headers = ["Nr.", "Status", "Ort in der ETS", "Aufgabe", "Erledigt", "Bemerkung"]
    for prio, rows in _numbered(worklist):
        if not rows:
            continue
        title, explanation = PRIORITIES[prio]
        excel.add_empty_row()
        excel.add_heading(f"{prio}  {title}  ({len(rows)}) – {explanation}", level=2)
        excel.add_table(headers, [
            [n, item.status, item.where, _task_text(item), XL_CHOICES[0], ""]
            for n, item in rows],
            center_cols=[1, 2, 5], wrap_cols=[3, 4, 6])

    ws = excel._current_sheet
    last = max(ws.max_row, 2)
    dv = DataValidation(type="list", formula1='"' + ",".join(XL_CHOICES) + '"',
                        allow_blank=True)
    ws.add_data_validation(dv)
    for r in range(1, last + 1):
        if ws.cell(r, 5).value in XL_CHOICES:
            dv.add(ws.cell(r, 5))
    green = PatternFill(start_color="DCEFD6", end_color="DCEFD6", fill_type="solid")
    grey = PatternFill(start_color="ECECEC", end_color="ECECEC", fill_type="solid")
    ws.conditional_formatting.add(f"A1:F{last}", FormulaRule(formula=['$E1="✓ erledigt"'], fill=green))
    ws.conditional_formatting.add(f"A1:F{last}", FormulaRule(formula=['$E1="– später"'], fill=grey))

    if worklist.done:
        excel.add_sheet("Erledigt")
        excel.set_column_widths([30, 70, 14])
        excel.add_heading("Erledigt (nach dem Re-Import festgestellt)", level=1)
        excel.add_table(["Ort in der ETS", "Aufgabe", "Erledigt am"],
                        [[i.where, i.task, i.done_on] for i in worklist.done],
                        wrap_cols=[2])
    excel.wb.active = 0
    excel.save(filepath)
    logger.info("ETS-Arbeitsliste Excel: %s", filepath)
    return worklist
