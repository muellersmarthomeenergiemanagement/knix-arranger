"""Inbetriebnahme-Checkliste: eine Zeile je Taste, Stockwerk-Gliederung."""
from knix_arranger.models.documentation import ChecklistItem
from knix_arranger.services.documentation_service import DocumentationService
from tests.test_multi_ga_check import _project


def test_eine_zeile_je_taste():
    project = _project()
    sections = DocumentationService(project)._checklist_sections()
    assert [(floor, room) for floor, room, _ in sections] == [("EG", "00 Halle")]
    title, rows = sections[0][2][0]
    assert title.startswith("1.1.51")
    tasten = [r[0] for r in rows]
    assert tasten[0] == "Gerät"
    # Taste 1 links: Befehl + Rückmeldung in EINER Zeile
    t1 = next(r for r in rows if r[0] == "1 links")
    assert "1/0/0" in t1[2] and "1/7/0" in t1[2]
    assert tasten.count("1 links") == 1


def test_ergebnis_je_zeile_zusammengefasst():
    saved = {
        ("r", "T", "1", "Taste 1", "1/0/0"): ChecklistItem(result="OK"),
        ("r", "T", "1", "Taste 1", "1/7/0"): ChecklistItem(result="Mangel", notes="LED dunkel"),
    }
    keys = list(saved)
    assert DocumentationService._combined_result(keys, saved) == ("Mangel", "LED dunkel")
    assert DocumentationService._combined_result(keys[:1], saved) == ("OK", "")
    assert DocumentationService._combined_result([("x",)], saved) == ("", "")


def test_abnahmeprotokoll_mit_profilen(tmp_path):
    import fitz
    from knix_arranger.models.company_profile import CompanyProfile
    from knix_arranger.models.client_profile import ClientProfile
    project = _project()
    project.client_profile = ClientProfile(name="Anita Muster", object_address="Chaletweg 5")
    company = CompanyProfile(company_name="Muster AG", user_name="Max Muster",
                             role="KNX-Systemintegrator")
    svc = DocumentationService(project, company_profile=company)
    protocol = svc.create_acceptance_protocol()
    assert protocol.integrator_name == "Max Muster, Muster AG"
    assert protocol.client_name == "Anita Muster"
    assert len(protocol.date) == 10 and protocol.date[2] == "."   # 01.10.2026
    path = str(tmp_path / "abnahme.pdf")
    svc.export_acceptance_protocol(path, protocol)
    text = "".join(pg.get_text() for pg in fitz.open(path))
    for expected in ("Umfang der Anlage", "Übergebene Unterlagen", "Mängelliste",
                     "Abnahmeentscheid", "Chaletweg 5", "Max Muster, Muster AG"):
        assert expected in text
    assert "Gepruefte" not in text


def test_excel_zeilenhoehe_nach_inhalt(tmp_path):
    import openpyxl
    path = str(tmp_path / "cl.xlsx")
    DocumentationService(_project()).export_checklists_excel(path)
    ws = openpyxl.load_workbook(path)["EG"]
    row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "1 links")
    lines = str(ws.cell(row, 3).value).count("\n") + 1
    assert lines >= 2
    assert ws.row_dimensions[row].height >= lines * 12


def test_excel_ausfuellen_und_einlesen(tmp_path):
    """Blätter je Stockwerk, Auswahlliste OK, Ergebnisse zurück in die App."""
    import openpyxl
    project = _project()
    path = str(tmp_path / "cl.xlsx")
    DocumentationService(project).export_checklists_excel(path)
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames[0] == "Übersicht" and "EG" in wb.sheetnames
    ws = wb["EG"]
    assert ws.column_dimensions["F"].hidden
    assert ws.data_validations.dataValidation
    row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "1 links")
    assert ws.cell(row, 4).value == "☐ offen"
    ws.cell(row, 4).value = "⚠ Mangel"
    ws.cell(row, 5).value = "LED dunkel"
    wb.save(path)

    stats = DocumentationService(project).import_checklists_excel(path)
    assert stats["items"] == 2 and stats["unknown"] == 0      # Befehl + Rückmeldung
    defects = [i for cl in project.checklists for i in cl.items if i.result == "Mangel"]
    assert {i.function_ga.split()[0] for i in defects} == {"1/0/0", "1/7/0"}
    assert all(i.notes == "LED dunkel" for i in defects)
    # «offen» überschreibt vorhandene Ergebnisse nicht
    DocumentationService(project).export_checklists_excel(path)
    wb = openpyxl.load_workbook(path)
    ws = wb["EG"]
    ws.cell(row, 4).value = "☐ offen"
    wb.save(path)
    DocumentationService(project).import_checklists_excel(path)
    assert all(i.result == "Mangel" for i in defects)
