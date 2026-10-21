"""Inbetriebnahme-Checkliste: Datum und Visum je Prüfpunkt abzeichnen –
in der Ansicht, in PDF/Excel, beim Einlesen und im Abnahmeprotokoll."""
from __future__ import annotations
import json
import os
import sys
from datetime import date, datetime

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pytest

from knix_arranger.models.company_profile import CompanyProfile
from knix_arranger.models.documentation import (
    ChecklistItem, CommissioningChecklist, RESULT_OK, RESULT_DEFECT, RESULT_OPEN,
    parse_check_date, display_check_date,
)
from knix_arranger.services.documentation_service import DocumentationService, _date_span
from tests.test_multi_ga_check import _project


# ── Modell ────────────────────────────────────────────────────────────────────

def test_set_result_zeichnet_ab_und_offen_leert():
    item = ChecklistItem()
    assert item.set_result(RESULT_OK, "MM", date(2026, 10, 14))
    assert (item.checked_date, item.checked_by) == ("2026-10-14", "MM")
    # gleiches Ergebnis: bisherige Abzeichnung bleibt
    assert not item.set_result(RESULT_OK, "RS", date(2026, 10, 15))
    assert (item.checked_date, item.checked_by) == ("2026-10-14", "MM")
    item.set_result(RESULT_DEFECT, "RS", date(2026, 10, 15))
    assert (item.checked_date, item.checked_by) == ("2026-10-15", "RS")
    item.set_result(RESULT_OPEN)
    assert (item.checked_date, item.checked_by) == ("", "")


def test_abzeichnung_wird_gespeichert():
    item = ChecklistItem(result=RESULT_OK, checked_date="2026-10-14", checked_by="MM")
    loaded = ChecklistItem.from_dict(item.to_dict())
    assert (loaded.checked_date, loaded.checked_by) == ("2026-10-14", "MM")
    # ältere Projektdateien ohne die Felder
    old = item.to_dict()
    del old["checked_date"], old["checked_by"]
    assert ChecklistItem.from_dict(old).checked_date == ""


@pytest.mark.parametrize("value, iso", [
    ("21.10.2026", "2026-10-21"), ("1.2.26", "2026-02-01"), ("2026-10-21", "2026-10-21"),
    (datetime(2026, 10, 21, 8, 30), "2026-10-21"), (date(2026, 10, 21), "2026-10-21"),
    ("31.02.2026", ""), ("morgen", ""), ("", ""), (None, ""),
])
def test_datum_lesen(value, iso):
    assert parse_check_date(value) == iso


def test_datum_anzeigen_und_zeitraum():
    assert display_check_date("2026-10-21") == "21.10.2026"
    assert display_check_date("") == ""
    assert _date_span("2026-10-14", "2026-10-14") == "14.10.2026"
    assert _date_span("2026-10-14", "2026-10-16") == "14.10.–16.10.2026"
    assert _date_span("2026-12-30", "2027-01-04") == "30.12.2026–04.01.2027"


def test_visum_aus_initialen_oder_namen():
    assert CompanyProfile(user_name="Max Muster").visum == "MM"
    assert CompanyProfile(user_name="Anna-Lena von Arx").visum == "ALVA"
    assert CompanyProfile(user_name="Max Muster", initials="mmu").visum == "mmu"
    assert CompanyProfile().visum == ""
    assert CompanyProfile.from_dict(CompanyProfile(initials="RS").to_dict()).initials == "RS"


# ── Ausdruck ──────────────────────────────────────────────────────────────────

def test_zeile_zeigt_spaetestes_datum_und_alle_pruefer():
    saved = {
        ("r", "T", "1", "Taste 1", "1/0/0"): ChecklistItem(
            result="OK", checked_date="2026-10-14", checked_by="MM"),
        ("r", "T", "1", "Taste 1", "1/7/0"): ChecklistItem(
            result="OK", checked_date="2026-10-16", checked_by="RS"),
        ("r", "T", "1", "Taste 1", "1/7/1"): ChecklistItem(),
    }
    assert DocumentationService._combined_check_mark(list(saved), saved) == (
        "2026-10-16", "MM, RS")
    assert DocumentationService._combined_check_mark([("x",)], saved) == ("", "")


def _first_key_items(project):
    """Prüfpunkte der Zeile "1 links" (Befehl + Rückmeldung)."""
    svc = DocumentationService(project)
    svc.init_project_checklists()
    _, _, cards = svc._checklist_sections()[0]
    keys = next(r[3] for r in cards[0][1] if r[0] == "1 links")
    by_key = svc._saved_results()
    return [by_key[k] for k in keys]


def test_pdf_mit_datum_und_visum(tmp_path):
    import fitz
    project = _project()
    for item in _first_key_items(project):
        item.set_result(RESULT_OK, "MM", date(2026, 10, 14))
    path = str(tmp_path / "cl.pdf")
    DocumentationService(project).export_checklists_pdf(path)
    text = "".join(pg.get_text() for pg in fitz.open(path))
    assert "Visum" in text and "14.10.2026" in text and "MM" in text


def test_excel_datum_visum_ausfuellen_und_einlesen(tmp_path):
    import openpyxl
    project = _project()
    items = _first_key_items(project)
    path = str(tmp_path / "cl.xlsx")
    DocumentationService(project).export_checklists_excel(path)
    wb = openpyxl.load_workbook(path)
    ws = wb["EG"]
    header = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Taste")
    assert [ws.cell(header, c).value for c in range(4, 8)] == ["OK", "Datum", "Visum", "Bemerkung"]
    row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "1 links")
    assert ws.cell(row, 5).number_format == "DD.MM.YYYY"
    ws.cell(row, 4).value = "✓ OK"
    ws.cell(row, 5).value = datetime(2026, 10, 15)
    ws.cell(row, 6).value = "RS"
    wb.save(path)

    DocumentationService(project).import_checklists_excel(path, visum="MM")
    assert {(i.result, i.checked_date, i.checked_by) for i in items} == {
        ("OK", "2026-10-15", "RS")}

    # Export zeigt die Abzeichnung wieder an
    DocumentationService(project).export_checklists_excel(path)
    ws = openpyxl.load_workbook(path)["EG"]
    assert ws.cell(row, 5).value.date() == date(2026, 10, 15)
    assert ws.cell(row, 6).value == "RS"


def test_excel_ohne_datum_visum_nimmt_dateidatum_und_einlesenden(tmp_path):
    import openpyxl
    project = _project()
    items = _first_key_items(project)
    path = str(tmp_path / "cl.xlsx")
    DocumentationService(project).export_checklists_excel(path)
    wb = openpyxl.load_workbook(path)
    ws = wb["EG"]
    row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "1 links")
    ws.cell(row, 4).value = "⚠ Mangel"
    wb.save(path)
    stamp = datetime(2026, 10, 12, 17, 0).timestamp()
    os.utime(path, (stamp, stamp))

    DocumentationService(project).import_checklists_excel(path, visum="MM")
    assert {(i.result, i.checked_date, i.checked_by) for i in items} == {
        ("Mangel", "2026-10-12", "MM")}


def test_excel_datum_visum_ohne_ergebnis(tmp_path):
    """Datum/Visum ausgefüllt, OK noch «offen»: Abzeichnung wird übernommen,
    das Ergebnis bleibt offen, die Zeile wird gemeldet."""
    import openpyxl
    project = _project()
    items = _first_key_items(project)
    path = str(tmp_path / "cl.xlsx")
    DocumentationService(project).export_checklists_excel(path)
    wb = openpyxl.load_workbook(path)
    ws = wb["EG"]
    row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "1 links")
    ws.cell(row, 5).value = datetime(2026, 10, 10)
    ws.cell(row, 6).value = "MM"
    wb.save(path)

    stats = DocumentationService(project).import_checklists_excel(path, visum="RS")
    assert stats["items"] == 2 and stats["unrated"] == 1
    assert {(i.result, i.checked_date, i.checked_by) for i in items} == {
        ("", "2026-10-10", "MM")}
    # erneuter Export zeigt die Abzeichnung wieder
    DocumentationService(project).export_checklists_excel(path)
    ws = openpyxl.load_workbook(path)["EG"]
    assert ws.cell(row, 6).value == "MM"


def test_excel_altes_format_bezug_in_spalte_f(tmp_path):
    """Bis 1.4.11 verteilte Checklisten (ohne Datum/Visum) bleiben einlesbar."""
    import openpyxl
    project = _project()
    items = _first_key_items(project)
    keys = [[i.room_id, i.be_type, i.be_number, i.check_type, i.function_ga] for i in items]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "EG"
    ws.append(["Taste", "Funktion", "Gruppenadressen", "OK", "Bemerkung", "Bezug"])
    ws.append(["1 links", "", "", "⚠ Mangel", "LED dunkel", json.dumps(keys)])
    path = str(tmp_path / "alt.xlsx")
    wb.save(path)

    stats = DocumentationService(project).import_checklists_excel(path, visum="MM")
    assert stats["items"] == 2
    assert all(i.result == "Mangel" and i.notes == "LED dunkel" and i.checked_by == "MM"
               for i in items)


# ── Abnahmeprotokoll ──────────────────────────────────────────────────────────

def test_abnahmeprotokoll_pruefung_je_raum(tmp_path):
    import fitz
    project = _project()
    project.checklists = [CommissioningChecklist(room_id="r1", room_name="00 Halle", items=[
        ChecklistItem(result="OK", checked_date="2026-10-14", checked_by="MM"),
        ChecklistItem(result="Mangel", checked_date="2026-10-16", checked_by="RS"),
        ChecklistItem(),
    ])]
    svc = DocumentationService(project)
    assert svc._checked_by_room() == [
        ["00 Halle", "3", "1", "1", "14.10.–16.10.2026", "MM, RS"]]
    path = str(tmp_path / "abnahme.pdf")
    svc.export_acceptance_protocol(path)
    text = "".join(pg.get_text() for pg in fitz.open(path))
    assert "Prüfung je Raum" in text and "MM, RS" in text


# ── Ansicht ───────────────────────────────────────────────────────────────────

@pytest.fixture(scope="module")
def qapp():
    from PySide6.QtWidgets import QApplication
    yield QApplication.instance() or QApplication([])


def _view(qapp, visum="MM"):
    from knix_arranger.models.project import KnxProject
    from knix_arranger.ui.views.commissioning_view import CommissioningView
    project = KnxProject(name="Test")
    cl = CommissioningChecklist(room_id="r1", room_name="M01 Musikzimmer", items=[
        ChecklistItem(check_type=f"Kanal {i + 1}", room_id="r1") for i in range(3)])
    project.checklists = [cl]
    view = CommissioningView(project, visum_provider=lambda: visum)
    view.set_project(project)
    view._tree.setCurrentItem(view._tree.topLevelItem(0))
    return view, cl


def test_ansicht_zeichnet_beim_setzen_ab(qapp):
    from knix_arranger.ui.views.commissioning_view import _COL_RES, _COL_DATE, _COL_VISUM
    view, cl = _view(qapp)
    combo = view._table.cellWidget(0, _COL_RES)
    combo.setCurrentIndex(combo.findData(RESULT_OK))
    today = date.today()
    assert (cl.items[0].checked_date, cl.items[0].checked_by) == (today.isoformat(), "MM")
    assert view._table.item(0, _COL_DATE).text() == today.strftime("%d.%m.%Y")
    assert view._table.item(0, _COL_VISUM).text() == "MM"
    combo.setCurrentIndex(combo.findData(RESULT_OPEN))
    assert view._table.item(0, _COL_DATE).text() == "" and cl.items[0].checked_by == ""


def test_ansicht_mehrere_zeilen_und_von_hand_aendern(qapp):
    from knix_arranger.ui.views.commissioning_view import _COL_DATE, _COL_VISUM
    view, cl = _view(qapp, visum="RS")
    view._apply_bulk([0, 1, 2], RESULT_OK)
    assert {i.checked_by for i in cl.items} == {"RS"}
    view._table.item(1, _COL_DATE).setText("3.10.26")
    view._table.item(1, _COL_VISUM).setText("MM")
    assert (cl.items[1].checked_date, cl.items[1].checked_by) == ("2026-10-03", "MM")
    assert view._table.item(1, _COL_DATE).text() == "03.10.2026"
