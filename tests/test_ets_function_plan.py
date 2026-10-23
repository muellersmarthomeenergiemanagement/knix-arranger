"""Funktionen planen (FA-530, FA-637): Gruppenbedienung mehrerer Elemente
mit eigenen Befehls-GAs (ohne Rückmeldung), z.B. "Fensterfront" für drei
Jalousien, auf Tasten legbar."""
import logging

import pytest

from knix_arranger.models.building import (
    Apartment, Areal, Bedienelement, Building, Floor, GewerkAssignment, Room,
    SensorFunktion, Wing,
)
from knix_arranger.models.ets_function import (
    EtsFunction, EtsFunctionRef, FunctionMember, KnixFunction,
)
from knix_arranger.models.project import KnxProject
from knix_arranger.services import ets_worklist
from knix_arranger.services.address_generator import regenerate_addresses
from knix_arranger.services.belegungsplan_service import BelegungsplanService
from knix_arranger.services.co_linking_service import CoLinkingService
from knix_arranger.services.ets_function_plan import (
    apply_function_buttons, check_plan, compare_plan, function_gas, member_links,
    restore_function_ids, sync_function_gas,
)
from knix_arranger.services.topology_engine import TopologyEngine
from knix_arranger.services.validation_engine import ValidationEngine


def _build(variant="A"):
    logging.disable(logging.CRITICAL)
    project = KnxProject(name="Funktionen")
    project.config.mg_variant = variant
    blinds = GewerkAssignment(gewerk_code="J", count=3,
                              element_labels=["Links", "Mitte", "Rechts"])
    lights = GewerkAssignment(gewerk_code="LD", count=2)
    wohnen = Room(number="01", name="Wohnen", gewerk_assignments=[
        blinds, lights, GewerkAssignment(gewerk_code="H", count=1)])
    kueche = Room(number="02", name="Küche", gewerk_assignments=[
        GewerkAssignment(gewerk_code="LD", count=1), GewerkAssignment(gewerk_code="H", count=1)])
    zone = Apartment(name="Wohnzone", rooms=[wohnen, kueche])
    floor = Floor(name="EG", short_code="EG", main_group_number=2, apartments=[zone])
    project.areal = Areal(buildings=[Building(name="H", wings=[Wing(floors=[floor])])])
    engine = TopologyEngine(project.config.topology_mode)
    project.topology = engine.calculate_topology(project.areal)
    engine.populate_devices(project.topology, project.all_rooms, project.gewerk_catalog)
    front = KnixFunction(name="Fensterfront", function_type="FT-7", room_id=wohnen.id,
                         members=[FunctionMember(blinds.id, n) for n in (1, 2, 3)])
    project.functions.append(front)
    regenerate_addresses(project)
    return project


@pytest.fixture
def project():
    yield _build()
    logging.disable(logging.NOTSET)


def _gas(project, function=None):
    function = function or project.functions[0]
    return {ga.function_name: ga for ga in function_gas(project.group_addresses)
            if ga.function_id == function.id}


def _room(project, name):
    return next(r for r in project.all_rooms if r.name == name)


def _channel_kos(project):
    """Je Jalousie ein Gerät mit einem KO je Funktion, wie aus der ETS: die
    GA des Elements und die Funktions-GA."""
    from knix_arranger.models.topology import CommunicationObject, Device
    line = project.topology.areas[0].lines[0]
    devices = {}
    for member, fga, own in member_links(project, project.functions[0]):
        nr = member.element_number
        device = devices.get(nr)
        if device is None:
            device = devices[nr] = Device(physical_address=f"1.1.{100 + nr}")
            line.devices.append(device)
        co = CommunicationObject(object_number=len(device.communication_objects),
                                 name=own.function_name, connected_gas=[own.address, fga.address])
        device.communication_objects.append(co)


def _cos_with(project, address):
    return [co for a in project.topology.areas for ln in a.lines for d in ln.devices
            for co in d.communication_objects if address in co.connected_gas]


def test_command_addresses_from_top_block(project):
    """F1 im obersten 5er-Block 250–254, Elemente darunter."""
    gas = _gas(project)
    assert {f: (ga.address, ga.datapoint_type) for f, ga in gas.items()} == {
        "AUF/AB": ("2/1/250", "DPST-1-8"), "STOPP": ("2/1/251", "DPST-1-7"),
        "POSITION HOEHE": ("2/1/252", "DPST-5-1"), "POSITION LAMELLEN": ("2/1/253", "DPST-5-1"),
        "SPERREN": ("2/1/254", "DPST-1-1")}
    assert gas["AUF/AB"].designation == "J_EG_01_F1 AUF/AB (Fensterfront)"
    element = [ga for ga in project.group_addresses.all_addresses()
               if ga.gewerk_code == "J" and not ga.function_id and not ga.is_placeholder]
    assert max(ga.sub_group for ga in element) < 200


def test_variant_b_same_addresses_without_findings():
    project = _build("B")
    gas = _gas(project)
    assert sorted(ga.address for ga in gas.values()) == [
        "2/1/250", "2/1/251", "2/1/252", "2/1/253", "2/1/254"]
    issues = ValidationEngine(project.gewerk_catalog).validate(
        project.group_addresses, project=project)
    addresses = {ga.address for ga in gas.values()}
    assert not [i for i in issues if i.address in addresses]


def test_stable_and_next_block(project):
    before = {ga.id: ga.address for ga in function_gas(project.group_addresses)}
    zone = project.areal.all_floors[0].apartments[0]
    lights = [a for r in project.all_rooms for a in r.gewerk_assignments if a.gewerk_code == "LD"]
    project.functions.append(KnixFunction(
        name="Licht Wohnzone", function_type="FT-6", zone_id=zone.id,
        members=[FunctionMember(a.id, 1) for a in lights]))
    project.functions.append(KnixFunction(
        name="Ganze Front", function_type="FT-7", room_id=_room(project, "Wohnen").id,
        number=2, members=[FunctionMember(project.functions[0].members[0].assignment_id, 1)]))
    regenerate_addresses(project)
    after = {ga.id: ga.address for ga in function_gas(project.group_addresses)}
    assert all(after[i] == a for i, a in before.items())
    assert _gas(project, project.functions[2])["AUF/AB"].address == "2/1/245"
    zone_gas = _gas(project, project.functions[1])
    assert sorted(zone_gas) == ["DIM", "E/A", "WERT"]
    assert zone_gas["E/A"].address == "2/0/250"
    assert zone_gas["E/A"].designation == "LD_EG_F1 E/A (Licht Wohnzone)"


def test_deleted_function_drops_gas(project):
    project.functions.clear()
    sync_function_gas(project)
    assert not function_gas(project.group_addresses)


def test_rows_on_all_member_channels(project):
    rows = BelegungsplanService().generate(project).actor_rows
    by_ga = {}
    for row in rows:
        by_ga.setdefault(row.ga_address, set()).add((row.physical_address, row.channel_number))
    assert all(len(by_ga[a]) == 3
               for a in ("2/1/250", "2/1/251", "2/1/252", "2/1/253", "2/1/254"))


def test_co_linking_proposals(project):
    proposals = CoLinkingService().generate_proposals(project)
    assert {p.ga_address for p in proposals} >= {"2/1/250", "2/1/251"}


def test_link_checks(project):
    _channel_kos(project)
    assert len(_cos_with(project, "2/1/250")) == 3
    assert not [i for i in check_plan(project) if i.rule_id == "FA-637"]
    _cos_with(project, "2/1/250")[1].connected_gas.remove("2/1/250")
    messages = [i.message for i in check_plan(project)]
    assert any("Funktions-GA 2/1/250 fehlt" in m for m in messages)
    project.topology.is_imported = True
    must = [i for i in ets_worklist.current_items(project) if i.rule_id == "FA-637"]
    assert len(must) == 1 and must[0].priority == ets_worklist.PRIO_MUST
    assert must[0].task == "2/1/250 verknüpfen (Funktion «Fensterfront»)."


def test_compare_with_ets(project):
    wohnen = _room(project, "Wohnen")
    assert not check_plan(project)                   # nichts eingelesen: keine Aufgabe
    wohnen.ets_functions.append(EtsFunction("Andere", "FT-1", []))
    issues = check_plan(project)
    assert [i.details["kind"] for i in issues] == ["anlegen"]
    assert issues[0].details["task"] == (
        "Funktion «Fensterfront» (Sonnenschutz) anlegen und zuordnen: "
        "Bewegen 2/1/250, Schritt/Stop 2/1/251.")
    wohnen.ets_functions.append(EtsFunction("Front", "FT-7", [
        EtsFunctionRef("MoveUpDown", "2/1/250"), EtsFunctionRef("StopStepUpDown", "12/1/3")]))
    status = compare_plan(project)[0]
    assert status.state == "ergänzen" and status.ets.name == "Front"
    assert [ga.address for ga in status.missing] == ["2/1/251"]
    project.topology.is_imported = True
    items = [i for i in ets_worklist.current_items(project) if i.rule_id == "FA-637"]
    assert [i.priority for i in items] == [ets_worklist.PRIO_FUNC]
    assert items[0].task == "Funktion «Front»: Schritt/Stop 2/1/251 zuordnen."


def test_button_on_function(project):
    wohnen = _room(project, "Wohnen")
    sf = SensorFunktion(function_id=project.functions[0].id)
    wohnen.bedienelemente.append(Bedienelement(
        element_type="Tastereinheit", channels=4, funktionen=[sf], is_auto=False))
    apply_function_buttons(project)
    assert sf.label == "Fensterfront"
    assert sf.ga_designation == "2/1/250  J_EG_01_F1 AUF/AB (Fensterfront)"
    assert [(e.ga_designation, e.role) for e in sf.extra_gas] == [
        ("2/1/251  J_EG_01_F1 STOPP (Fensterfront)", "befehl"),
        ("2/1/252  J_EG_01_F1 POSITION HOEHE (Fensterfront)", "befehl"),
        ("2/1/253  J_EG_01_F1 POSITION LAMELLEN (Fensterfront)", "befehl"),
        ("2/1/254  J_EG_01_F1 SPERREN (Fensterfront)", "befehl")]
    # Umbenennen wird nachgeführt
    project.functions[0].name = "Front Süd"
    sync_function_gas(project)
    apply_function_buttons(project)
    assert sf.ga_designation == "2/1/250  J_EG_01_F1 AUF/AB (Front Süd)"
    from knix_arranger.services.sensor_service import refresh_bedienelemente
    refresh_bedienelemente(project)
    be = next(b for b in wohnen.bedienelemente if b.funktionen and b.funktionen[0] is sf)
    gas = {fa.function_ga.split()[0] for fa in be.function_assignments}
    assert gas >= {"2/1/250", "2/1/251"}
    # Funktion gelöscht: Taste bleibt mit Namen, ohne GA
    project.functions.clear()
    sync_function_gas(project)
    apply_function_buttons(project)
    assert sf.ga_designation == "" and sf.extra_gas == [] and sf.label == "Front Süd"
    again = SensorFunktion.from_dict(sf.to_dict())
    assert again.function_id == sf.function_id


def test_restore_after_reimport(project):
    for ga in function_gas(project.group_addresses):
        ga.function_id = ""
    assert restore_function_ids(project) == 5
    assert len(_gas(project)) == 5


def test_renumbering_keeps_function_gas(project):
    from knix_arranger.services.renumber_service import plan_renumbering
    plan = plan_renumbering(project)
    gas = [ga for ga in plan.structure.all_addresses() if ga.function_id]
    assert sorted(ga.address for ga in gas) == [
        "2/1/250", "2/1/251", "2/1/252", "2/1/253", "2/1/254"]
    assert not [c for c in plan.changes if "_F1" in c.new_designation]


def test_roundtrip(project):
    again = KnxProject.from_dict(project.to_dict())
    assert again.functions == project.functions
    assert len([ga for ga in again.group_addresses.all_addresses() if ga.function_id]) == 5


def test_removed_member_reported(project):
    project.functions[0].members.append(FunctionMember("weg", 1))
    assert any("entfernte(s) Element(e)" in i.message for i in check_plan(project))


def test_view_dialog_and_button_choice(project):
    import os
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtCore import Qt
    from PySide6.QtWidgets import QApplication
    from knix_arranger.ui.views.ets_functions_view import EtsFunctionsView
    from knix_arranger.ui.dialogs.function_edit_dialog import FunctionEditDialog
    _app = QApplication.instance() or QApplication([])
    view = EtsFunctionsView()
    view.set_project(project)
    function_item = view._plan_tree.topLevelItem(0).child(0)
    assert function_item.text(0) == "Fensterfront (Sonnenschutz)"
    assert function_item.childCount() == 2
    gas_node, members_node = function_item.child(0), function_item.child(1)
    assert gas_node.text(0) == "Funktions-GAs (5)" and gas_node.childCount() == 5
    assert members_node.text(0) == "Bediente Elemente (3)"
    first = members_node.child(0)
    assert [first.text(c) for c in range(4)] == ["", "J_EG_01_01 (Links)", "", ""]
    assert first.text(4).startswith("Aktor ") and ", Kanal " in first.text(4)

    dialog = FunctionEditDialog(project, project.functions[0])
    assert dialog._members.count() == 3
    assert all(dialog._members.item(i).checkState() == Qt.Checked for i in range(3))
    dialog._members.item(2).setCheckState(Qt.Unchecked)
    dialog._accept()
    edited = dialog.result_function()
    assert len(edited.members) == 2 and edited.number == 1

    from knix_arranger.services.bauherr_form_service import BauherrFormService
    from knix_arranger.ui.views import bauherr_form_view as bfv
    wohnen = _room(project, "Wohnen")
    be = Bedienelement(element_type="Tastereinheit", channels=4, is_auto=False)
    wohnen.bedienelemente.append(be)
    slot = bfv._SlotWidget(be, None, BauherrFormService(project), "Taste 1", wohnen)
    index = next(i for i in range(slot._combo.count())
                 if slot._combo.itemText(i).startswith("Funktion: Fensterfront – Sonnenschutz"))
    slot._combo.setCurrentIndex(index)
    sf = be.funktionen[0]
    assert sf.function_id == project.functions[0].id
    assert sf.ga_designation.startswith("2/1/250")
    assert slot._select_current(sf)
    assert slot._combo.currentIndex() == index


def test_bauherr_form_offers_and_reads_functions(project, tmp_path):
    from openpyxl import load_workbook
    from knix_arranger.services.bauherr_form_service import (
        BauherrFormService, CHOICES_SHEET, FUNCTION_PREFIX)
    wohnen = _room(project, "Wohnen")
    be = Bedienelement(element_type="Tastereinheit", channels=4, is_auto=False,
                       funktionen=[SensorFunktion(gewerk_code="LD", element_number=1)])
    wohnen.bedienelemente.append(be)
    path = tmp_path / "form.xlsx"
    service = BauherrFormService(project)
    service.generate_form(str(path))
    wb = load_workbook(path)
    formulas = [dv.formula1 for ws in wb.worksheets
                for dv in ws.data_validations.dataValidation]
    assert formulas and all(f.startswith(f"'{CHOICES_SHEET}'!$") for f in formulas)
    choices = wb[CHOICES_SHEET]
    column = [c.value for c in next(choices.iter_cols(min_col=1, max_col=1))]
    assert column[1] == "Funktion: Fensterfront" and "Garagentor" in column
    assert "Sonstiges" in column
    # Bauherr wählt die Funktion für Taste 2
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row[:2]:
                if str(cell.value or "").startswith("T2  "):
                    cell.value = "T2  " + FUNCTION_PREFIX + "Fensterfront"
    wb.save(path)
    assert service.import_form(str(path)) >= 1
    sf = next(f for f in be.funktionen if f.function_id)
    assert sf.function_id == project.functions[0].id
    assert sf.ga_designation.startswith("2/1/250")


def test_heating_operating_mode_for_zone(project):
    zone = project.areal.all_floors[0].apartments[0]
    heating = [a for r in project.all_rooms for a in r.gewerk_assignments if a.gewerk_code == "H"]
    project.functions.append(KnixFunction(
        name="Wohnzone Abwesend", function_type="FT-9", zone_id=zone.id,
        members=[FunctionMember(a.id, 1) for a in heating]))
    sync_function_gas(project)
    gas = _gas(project, project.functions[-1])
    assert list(gas) == ["UMSCHALTEN BETRIEBSART"]
    ga = gas["UMSCHALTEN BETRIEBSART"]
    assert (ga.address, ga.datapoint_type) == ("2/2/250", "DPST-20-102")
    assert ga.designation == "H_EG_F1 UMSCHALTEN BETRIEBSART (Wohnzone Abwesend)"
    rows = BelegungsplanService().generate(project).actor_rows
    assert len({(r.physical_address, r.channel_number) for r in rows
                if r.ga_address == "2/2/250"}) == 2
    _room(project, "Wohnen").ets_functions.append(EtsFunction("x", "FT-1", []))
    task = next(i for i in check_plan(project)
                if "Wohnzone Abwesend" in i.message).details["task"]
    assert task.endswith("anlegen und zuordnen: Betriebsart 2/2/250.")


def test_element_labels_for_topology_channels():
    """Kanal der Topologie: bedientes Element mit Stockwerk; importierte GA
    mit Elementnummer nur in der Bezeichnung, Sammel-GA ohne Element."""
    from knix_arranger.models.building import (
        Apartment, Building, Floor, GewerkAssignment, Room, Wing,
    )
    from knix_arranger.models.group_address import (
        GroupAddress, GroupAddressStructure, MainGroup, MiddleGroup,
    )
    from knix_arranger.models.project import KnxProject
    from knix_arranger.services.ets_function_plan import element_labels_by_address
    project = KnxProject(name="Kanal")
    assignment = GewerkAssignment(gewerk_code="J", count=4)
    room = Room(number="04", name="Schlafen", gewerk_assignments=[assignment])
    floor = Floor(name="Dachgeschoss", short_code="DG",
                  apartments=[Apartment(rooms=[room])])
    project.areal.buildings.append(Building(wings=[Wing(floors=[floor])]))
    gas = [GroupAddress(main_group=4, middle_group=1, sub_group=n, designation=d,
                        assignment_id=assignment.id)
           for n, d in ((40, "J.DG.04.01_move  ( Berg )"),
                        (45, "J.DG.04.02+J.DG.04.03_move"))]
    gas.append(GroupAddress(main_group=4, middle_group=1, sub_group=60,
                            designation="J_DG_04_03 AUF/AB (Schlafen)",
                            assignment_id=assignment.id, element_number=3))
    project.group_addresses = GroupAddressStructure(main_groups=[MainGroup(
        number=4, middle_groups=[MiddleGroup(number=1, group_addresses=gas)])])
    assert element_labels_by_address(project) == {
        "4/1/40": "J_DG_04_01 (Schlafen)", "4/1/60": "J_DG_04_03 (Schlafen)"}


def _imported_dg():
    """Importiertes Projekt wie das Chalet: DG in HG 4 «DG» (nicht HG 7),
    Element-GAs nur mit Zuweisung, Nummer und Funktion in der Bezeichnung."""
    from knix_arranger.models.building import Apartment, Building, Floor, Room, Wing
    from knix_arranger.models.ets_function import FunctionMember, KnixFunction
    from knix_arranger.models.group_address import (
        GroupAddress, GroupAddressStructure, MainGroup, MiddleGroup,
    )
    from knix_arranger.models.topology import Topology
    project = KnxProject(name="Import")
    project.topology = Topology(is_imported=True)
    assignment = GewerkAssignment(gewerk_code="J", count=2)
    room = Room(number="04", name="Schlafen", gewerk_assignments=[assignment])
    floor = Floor(name="Dachgeschoss", short_code="DG", main_group_number=7,
                  apartments=[Apartment(rooms=[room])])
    project.areal.buildings.append(Building(wings=[Wing(floors=[floor])]))

    def ga(mg, sub, designation):
        return GroupAddress(main_group=4, middle_group=mg, sub_group=sub,
                            designation=designation, assignment_id=assignment.id)
    commands = [ga(1, 40 + i, f"J.DG.04.0{n}_{s}")
                for n in (1, 2) for i, s in enumerate(
                    ("move", "step", "hoehe", "winkel"), start=(n - 1) * 5)]
    feedback = [ga(6, 42, "J.DG.04.01_hoehe")]
    project.group_addresses = GroupAddressStructure(main_groups=[MainGroup(
        number=4, name="DG", middle_groups=[
            MiddleGroup(number=6, group_addresses=feedback),
            MiddleGroup(number=1, group_addresses=commands)])])
    project.functions.append(KnixFunction(
        name="Front", function_type="FT-7", room_id=room.id, number=1,
        members=[FunctionMember(assignment.id, 1), FunctionMember(assignment.id, 2)]))
    return project


def test_imported_function_gas_in_floor_main_group():
    from knix_arranger.services.ets_function_plan import (
        floor_main_groups, function_gas, member_links, sync_function_gas,
    )
    project = _imported_dg()
    assert list(floor_main_groups(project).values()) == [4]
    sync_function_gas(project)
    gas = {g.function_name: g for g in function_gas(project.group_addresses)}
    assert gas["AUF/AB"].address == "4/1/250"
    assert gas["AUF/AB"].designation == "J_DG_04_F1 AUF/AB (Front)"
    links = {(m.element_number, f.function_name): own.address if own else None
             for m, f, own in member_links(project, project.functions[0])}
    assert links[(1, "AUF/AB")] == "4/1/40"
    assert links[(1, "POSITION HOEHE")] == "4/1/42"   # Befehl, nicht Rückmeldung 4/6/42
    assert links[(2, "POSITION LAMELLEN")] == "4/1/48"
    assert links[(1, "SPERREN")] is None


def test_function_ga_from_ets_keeps_address_and_name():
    from knix_arranger.models.group_address import GroupAddress
    from knix_arranger.services.ets_function_plan import (
        function_gas, restore_function_ids, sync_function_gas,
    )
    project = _imported_dg()
    mg = project.group_addresses.main_groups[0].middle_groups[1]
    mg.group_addresses.append(GroupAddress(
        main_group=4, middle_group=1, sub_group=120, designation="J_04_F1 AUF/AB (Front)"))
    assert restore_function_ids(project) == 1
    sync_function_gas(project)
    gas = {g.function_name: g for g in function_gas(project.group_addresses)}
    assert (gas["AUF/AB"].address, gas["AUF/AB"].designation) == (
        "4/1/120", "J_04_F1 AUF/AB (Front)")
    assert gas["AUF/AB"].ets_known and not gas["STOPP"].ets_known
    assert gas["STOPP"].address == "4/1/251"
    assert GroupAddress.from_dict(gas["AUF/AB"].to_dict()).ets_known


def test_function_on_button_operates_members_and_shows_in_room_book(project):
    """Projekt_23: Jalousien nur über die Funktion bedient galten als «ohne
    Bedienung», die Taste hiess nach der Funktion statt «Taste n», und die
    Funktion fehlte im Raumbuch."""
    from knix_arranger.services.operation_check import unoperated_elements
    from knix_arranger.services.report_service import ReportService
    from knix_arranger.services.room_sheet import build_room_sheets
    from knix_arranger.services.sensor_service import refresh_bedienelemente
    wohnen = _room(project, "Wohnen")
    blinds = next(a for a in wohnen.gewerk_assignments if a.gewerk_code == "J")
    open_blinds = {(u.room_id, u.element_number) for u in unoperated_elements(project)
                   if u.gewerk_code == "J"}
    assert open_blinds == {(wohnen.id, n) for n in (1, 2, 3)}

    wohnen.bedienelemente.append(Bedienelement(
        element_type="Tastereinheit", channels=4, is_auto=False, participant_number="1.1.101",
        funktionen=[SensorFunktion(gewerk_code="LD", element_number=1),
                    SensorFunktion(function_id=project.functions[0].id)]))
    apply_function_buttons(project)
    refresh_bedienelemente(project)
    assert not [u for u in unoperated_elements(project) if u.gewerk_code == "J"]

    groups, _missing = ReportService(project)._room_gewerk_groups()
    sheet = next(s for s in build_room_sheets(
        project, groups, project.gewerk_catalog, BelegungsplanService().generate(project))
        if s.room is wohnen)
    function_row = next(e for e in sheet.elements if e.category == "funktion")
    assert function_row.label == "Fensterfront (J 01, 02, 03)"
    assert function_row.operated_by == ["1.1.101 Taste 2"]
    assert function_row.ga_ranges == "2/1/250–254"
    blind_rows = [e for e in sheet.elements if e.gewerk_code == "J"]
    assert [e.operated_by for e in blind_rows] == [["1.1.101 Taste 2 · Funktion Fensterfront"]] * 3
    assert all("2/1/250" not in e.ga_ranges for e in blind_rows)
    assert blinds.count == 3


def test_existing_function_at_200_keeps_its_address(project):
    """Vor 1.4.18 lagen Funktionen ab UG 200: sie behalten ihre Adresse (sonst
    stimmte die ETS nicht mehr), nur neue kommen von oben."""
    for ga in function_gas(project.group_addresses):
        ga.sub_group -= 50                      # wie früher: 2/1/200–204
    project.group_addresses.main_groups.sort(key=lambda h: h.number)
    regenerate_addresses(project)
    assert sorted(ga.address for ga in _gas(project).values()) == [
        "2/1/200", "2/1/201", "2/1/202", "2/1/203", "2/1/204"]
    project.functions.append(KnixFunction(
        name="Zweite", function_type="FT-7", room_id=_room(project, "Wohnen").id, number=2,
        members=[FunctionMember(project.functions[0].members[0].assignment_id, 1)]))
    regenerate_addresses(project)
    assert _gas(project, project.functions[1])["AUF/AB"].address == "2/1/250"


# ── GAs je Funktion anpassen (User 2026-10-23) ────────────────────────────

def test_omit_and_custom_points_share_the_block(project):
    from knix_arranger.models.ets_function import CustomPoint
    front = project.functions[0]
    front.omitted = ["POSITION LAMELLEN"]
    front.custom_points = [CustomPoint("SZENE", "DPST-17-1")]
    sync_function_gas(project)
    gas = {f: ga.address for f, ga in _gas(project).items()}
    # SZENE nimmt den Platz der weggelassenen Lamellen
    assert gas == {"AUF/AB": "2/1/250", "STOPP": "2/1/251", "POSITION HOEHE": "2/1/252",
                   "SZENE": "2/1/253", "SPERREN": "2/1/254"}
    assert _gas(project)["SZENE"].datapoint_type == "DPST-17-1"
    # mehr eigene Punkte: zweiter Block, beide Blöcke zusammen von oben
    front.custom_points += [CustomPoint("BESCHATTUNG", "DPST-1-1"),
                            CustomPoint("WIND", "DPST-1-5")]
    sync_function_gas(project)
    gas = {f: ga.address for f, ga in _gas(project).items()}
    assert gas["BESCHATTUNG"] == "2/1/250" and gas["WIND"] == "2/1/251"
    assert gas["AUF/AB"] == "2/1/245"
    again = KnixFunction.from_dict(front.to_dict())
    assert (again.omitted, [c.function for c in again.custom_points]) == (
        ["POSITION LAMELLEN"], ["SZENE", "BESCHATTUNG", "WIND"])


def test_linked_existing_ga_is_used_not_created(project):
    from knix_arranger.services.ets_function_plan import function_ga_map
    front = project.functions[0]
    lock = next(ga for ga in project.group_addresses.all_addresses()
                if ga.function_name == "SPERREN" and not ga.function_id)
    front.linked_gas = {"SPERREN": f"{lock.address}  {lock.designation}"}
    sync_function_gas(project)
    assert "SPERREN" not in _gas(project)                     # nicht angelegt
    mapping = function_ga_map(project, front)
    assert mapping["SPERREN"] is lock and lock.function_id == ""   # bleibt Element-GA
    wohnen = _room(project, "Wohnen")
    sf = SensorFunktion(function_id=front.id)
    wohnen.bedienelemente.append(Bedienelement(
        element_type="Tastereinheit", channels=4, funktionen=[sf], is_auto=False))
    apply_function_buttons(project)
    assert any(e.ga_designation.startswith(lock.address) for e in sf.extra_gas)


def test_user_defined_function_without_members(project):
    from knix_arranger.models.ets_function import CustomPoint
    from knix_arranger.services.ets_function_plan import (
        GENERAL_MIDDLE_GROUP, check_plan, function_ga_map,
    )
    party = KnixFunction(name="Party", function_type="FT-0",
                         room_id=_room(project, "Wohnen").id, number=2,
                         custom_points=[CustomPoint("SZENE", "DPST-17-1")])
    project.functions.append(party)
    sync_function_gas(project)
    [(point, ga)] = function_ga_map(project, party).items()
    assert (point, ga.address, ga.middle_group) == ("SZENE", "2/4/250", GENERAL_MIDDLE_GROUP)
    assert ga.designation == "SF_EG_01_F2 SZENE (Party)"
    assert not [i for i in check_plan(project) if "Party" in i.message]


def test_custom_point_without_element_ga_needs_manual_link(project):
    from knix_arranger.models.ets_function import CustomPoint
    from knix_arranger.services.ets_function_plan import check_plan
    project.functions[0].custom_points = [CustomPoint("SZENE", "DPST-17-1")]
    sync_function_gas(project)
    issues = [i for i in check_plan(project) if "von Hand" in i.message]
    assert len(issues) == 1 and "SZENE" in issues[0].message
    assert issues[0].details["kind"] == "verknüpfen-hand"


def test_adoption_before_1_4_18_plans_no_points():
    data = KnixFunction(name="Licht Zentral", function_type="FT-1",
                        ets_gas={"E/A": "3/4/89  Raum5_Licht Zentral"}).to_dict()
    data.pop("omitted", None)
    old = KnixFunction.from_dict(data)
    from knix_arranger.services.ets_function_plan import function_points, planned_points
    assert old.omitted is None
    assert [p.function for p in function_points(old)] == ["E/A"]
    assert planned_points(old) == []


def test_dialog_edits_gas(project):
    import os
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtWidgets import QApplication
    from knix_arranger.models.ets_function import CustomPoint
    from knix_arranger.ui.dialogs.function_edit_dialog import FunctionEditDialog
    _app = QApplication.instance() or QApplication([])
    dialog = FunctionEditDialog(project, project.functions[0])
    assert dialog._gas.rowCount() == 5
    assert dialog._gas.item(0, 1).text().startswith("2/1/250")
    dialog._on_kind("POSITION LAMELLEN", False, "omit")
    dialog._custom.append(CustomPoint("SZENE", "DPST-17-1"))
    dialog._fill_gas()
    assert dialog._gas.rowCount() == 6
    assert dialog._gas.item(3, 1).text() == "– weggelassen"
    dialog._accept()
    edited = dialog.result_function()
    assert edited.omitted == ["POSITION LAMELLEN"]
    assert [c.function for c in edited.custom_points] == ["SZENE"]
